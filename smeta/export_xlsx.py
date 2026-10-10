"""XLSX export: the customer's estimate (company header, INN, number) and the internal margin file."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import MATERIAL, Estimate, Totals
from .money import money
from .profit import ProfitReport

MONEY_FMT = "# ##0.00"
QTY_FMT = "# ##0.###"
THIN = Side(style="thin", color="FF808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="FFEFEFEF")
TOTAL_FILL = PatternFill("solid", fgColor="FFF7F7F7")
WARN_FILL = PatternFill("solid", fgColor="FFFFF0F0")

COLUMNS = [
    ("№", 6),
    ("Обоснование", 16),
    ("Наименование работ и затрат", 52),
    ("Ед. изм.", 10),
    ("Кол-во", 11),
    ("Цена за ед., ₽", 15),
    ("Всего, ₽", 16),
    ("Примечание", 26),
]


def _write_row(ws, row: int, values: list, *, bold=False, fill=None, border=True, wrap_col=3):
    for index, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=index, value=value)
        cell.font = Font(name="Calibri", size=10, bold=bold)
        if border:
            cell.border = BOX
        if fill:
            cell.fill = fill
        if index == wrap_col:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        elif index in (1, 4):
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif index >= 5:
            cell.alignment = Alignment(horizontal="right", vertical="top")
        if index == 5:
            cell.number_format = QTY_FMT
        elif index in (6, 7):
            cell.number_format = MONEY_FMT
    return row + 1


def _title_row(ws, row: int, text: str, *, size=12, bold=True, align="center") -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name="Calibri", size=size, bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    return row + 1


def _totals_block(ws, row: int, totals: Totals, estimate: Estimate) -> int:
    def line(label: str, value, *, bold=False):
        nonlocal row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = Font(name="Calibri", size=10, bold=bold)
        label_cell.alignment = Alignment(horizontal="right")
        label_cell.border = BOX
        value_cell = ws.cell(row=row, column=7, value=float(money(value)))
        value_cell.font = Font(name="Calibri", size=10, bold=bold)
        value_cell.number_format = MONEY_FMT
        value_cell.border = BOX
        value_cell.alignment = Alignment(horizontal="right")
        value_cell.fill = TOTAL_FILL
        ws.cell(row=row, column=8).border = BOX
        row += 1

    line("Прямые затраты, всего", totals.direct, bold=True)
    if totals.ot:
        line("в том числе оплата труда рабочих", totals.ot)
    if totals.em:
        line("эксплуатация машин и механизмов", totals.em)
    if totals.mat:
        line("материалы и изделия", totals.mat)
    if totals.other:
        line("работы и затраты по договорным ценам", totals.other)
    if totals.nr:
        line("Накладные расходы", totals.nr)
    if totals.sp:
        line("Сметная прибыль", totals.sp)
    line("Итого без НДС", totals.subtotal, bold=True)
    if estimate.vat_pct:
        line(f"НДС {totals.vat_pct:g}%", totals.vat)
    line("ВСЕГО по смете", totals.total, bold=True)
    if totals.tz:
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=f"Справочно: затраты труда {float(totals.tz):.1f} чел.-ч")
        cell.font = Font(name="Calibri", size=9, italic=True)
        row += 1
    return row


def estimate_to_xlsx(estimate: Estimate, path: str | Path) -> Path:
    """Customer-facing estimate. Purchase prices, discounts and margin are never written here."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for index, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    row = 1
    for line in estimate.company.header_lines():
        row = _title_row(ws, row, line, size=10, bold=line == estimate.company.name, align="left")
    row += 1
    number = f" № {estimate.number}" if estimate.number else ""
    row = _title_row(ws, row, f"ЛОКАЛЬНАЯ СМЕТА{number}", size=14)
    if estimate.title:
        row = _title_row(ws, row, estimate.title, size=11, bold=False)
    row += 1
    for label, value in (
        ("Заказчик", estimate.customer),
        ("Объект", estimate.object_name),
        ("Основание", estimate.basis),
        ("Уровень цен", f"{estimate.indices.region} {estimate.indices.period}".strip()),
        ("Дата", estimate.date.strftime("%d.%m.%Y")),
    ):
        if value:
            row = _title_row(ws, row, f"{label}: {value}", size=10, bold=False, align="left")
    row += 1

    row = _write_row(ws, row, [title for title, _ in COLUMNS], bold=True, fill=HEAD_FILL)
    header_row = row - 1
    ws.freeze_panes = ws.cell(row=row, column=1)

    number_counter = 0
    sections = estimate.sections()
    for section_name, positions in sections:
        if section_name and len(sections) > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
            cell = ws.cell(row=row, column=1, value=section_name)
            cell.font = Font(name="Calibri", size=10, bold=True)
            cell.fill = HEAD_FILL
            cell.border = BOX
            row += 1
        for position in positions:
            number_counter += 1
            note = position.note
            if position.kind == MATERIAL and "материал" not in note.lower():
                note = (note + " " if note else "") + "материал"
            row = _write_row(
                ws,
                row,
                [
                    number_counter,
                    position.code or "договорная цена",
                    position.name,
                    position.unit,
                    float(position.qty),
                    float(position.unit_price_now(estimate.indices)),
                    float(position.total(estimate.indices)),
                    note,
                ],
            )
    ws.auto_filter.ref = f"A{header_row}:H{row - 1}"

    row += 1
    row = _totals_block(ws, row, estimate.totals(), estimate)

    row += 2
    signer = estimate.company.signer or ""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(row=row, column=1, value=f"Составил: {estimate.company.signer_title} _______________ {signer}").font = Font(size=10)
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=8)
    ws.cell(row=row, column=6, value="Заказчик: _______________").font = Font(size=10)

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def margin_to_xlsx(report: ProfitReport, path: str | Path, estimate: Estimate | None = None) -> Path:
    """Internal file: sell price vs cost per line. Marked so it cannot be confused with the estimate."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Профит"
    widths = [6, 46, 10, 11, 15, 15, 15, 15, 10, 26]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(widths))
    warning = ws.cell(row=1, column=1, value="ДЛЯ ВНУТРЕННЕГО ИСПОЛЬЗОВАНИЯ — НЕ ОТПРАВЛЯТЬ ЗАКАЗЧИКУ")
    warning.font = Font(size=12, bold=True, color="FFB00020")
    warning.fill = WARN_FILL
    warning.alignment = Alignment(horizontal="center")
    row = 2
    if estimate is not None:
        label = f"Смета № {estimate.number or '—'} · {estimate.customer or ''} · {estimate.object_name or ''}"
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(widths))
        ws.cell(row=row, column=1, value=label).font = Font(size=10, italic=True)
        row += 2

    headers = ["№", "Позиция", "Ед.", "Кол-во", "В смете, ₽", "Себестоимость, ₽",
               "Маржа, ₽", "Маржа, %", "Оплата", "Источник"]
    for index, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=index, value=title)
        cell.font = Font(size=10, bold=True)
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    row += 1

    counter = 0
    for block_title, rows in (("Работы", report.works), ("Материалы", report.materials)):
        if not rows:
            continue
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(widths))
        cell = ws.cell(row=row, column=1, value=block_title)
        cell.font = Font(size=10, bold=True)
        cell.fill = HEAD_FILL
        row += 1
        for item in rows:
            counter += 1
            values = [
                counter,
                item.name,
                item.unit,
                float(item.qty),
                float(item.sell),
                None if item.cost is None else float(item.cost),
                None if item.margin is None else float(item.margin),
                None if item.margin_pct is None else float(item.margin_pct),
                item.payment if item.cost is not None else "",
                item.source or ("себестоимость не известна" if item.cost is None else ""),
            ]
            for index, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=index, value=value)
                cell.font = Font(size=10, bold=False, color="FFB00020" if item.loss else "FF000000")
                cell.border = BOX
                if index in (5, 6, 7):
                    cell.number_format = MONEY_FMT
                if index == 4:
                    cell.number_format = QTY_FMT
                if item.loss:
                    cell.fill = WARN_FILL
            row += 1

    row += 1
    for label, value, bold in (
        ("Сумма по смете", report.sell_total, False),
        ("Себестоимость", report.cost_total, False),
        ("Маржа работ", report.works_margin, False),
        ("Экономия на материалах", report.materials_margin, False),
        ("ИТОГО в карман", report.margin_total, True),
    ):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = Font(size=10, bold=bold)
        label_cell.alignment = Alignment(horizontal="right")
        value_cell = ws.cell(row=row, column=7, value=float(money(value)))
        value_cell.font = Font(size=10, bold=bold)
        value_cell.number_format = MONEY_FMT
        value_cell.fill = TOTAL_FILL
        row += 1
    payments = report.by_payment
    if payments:
        ws.cell(row=row, column=2, value="Закупка: " + ", ".join(f"{k} {float(v):,.2f} ₽" for k, v in payments.items())).font = Font(size=10, italic=True)
        row += 1

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
