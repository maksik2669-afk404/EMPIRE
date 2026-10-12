"""Importer for official data: FSNB-2022 (FER/GESN) rates and FGIS CS price indices.

The state portal publishes this data as XLSX, XML, DBF or ZIP archives of those, and the
column wording changes from release to release. So nothing here is hard-wired to one
layout: the importer finds the header row itself, classifies every column by meaning, and
returns a report saying what it mapped, what it skipped and why. A silent import is the
one thing worse than a failed one.

    python3 -m smeta.importer ФСНБ-2022_Часть1.xlsx --kind rates --out data/rates/fer2022.csv
    python3 -m smeta.importer индексы_IV_кв_2026.xlsx --kind indices --out data/rates/indices.csv
"""
from __future__ import annotations

import csv
import io
import struct
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree

from .catalog import Indices, Rate
from .money import ONE, ZERO, dec

RATE_FIELDS = ("code", "name", "unit", "ot", "em", "zpm", "mat", "tz", "tzm", "pz",
               "work_type", "source", "keywords")
INDEX_FIELDS = ("name", "region", "period", "group", "ot", "em", "mat", "zpm", "smr")
MAX_HEADER_SCAN = 40


from .headers import (  # noqa: F401
    GENERIC_HEADERS, classify, match_columns, normalize_header, sbornik_work_type,
)

# ---------------------------------------------------------------- report


@dataclass
class ImportReport:
    kind: str = "rates"
    source: str = ""
    fmt: str = ""
    sheet: str = ""
    header_row: int = 0
    columns: dict[str, str] = field(default_factory=dict)
    rows_seen: int = 0
    imported: int = 0
    skipped: Counter = field(default_factory=Counter)
    warnings: list[str] = field(default_factory=list)
    samples: list[str] = field(default_factory=list)

    def merge(self, other: "ImportReport") -> "ImportReport":
        self.rows_seen += other.rows_seen
        self.imported += other.imported
        self.skipped.update(other.skipped)
        for warning in other.warnings:
            if warning not in self.warnings:
                self.warnings.append(warning)
        self.samples = (self.samples + other.samples)[:3]
        if not self.columns:
            self.columns = other.columns
            self.sheet, self.header_row, self.fmt = other.sheet, other.header_row, other.fmt
        return self

    def as_text(self) -> str:
        what = "расценок" if self.kind == "rates" else "индексов"
        lines = [f"Импорт {what}: {self.imported} из {self.rows_seen} строк"]
        if self.fmt:
            place = f"{self.fmt}" + (f", лист «{self.sheet}»" if self.sheet else "")
            lines.append(f"Файл: {self.source or '—'} ({place}, заголовок в строке {self.header_row})")
        if self.columns:
            lines.append("Распознаны столбцы: " + ", ".join(f"{k} ← «{v}»" for k, v in self.columns.items()))
        for reason, count in self.skipped.most_common():
            lines.append(f"Пропущено {count}: {reason}")
        for warning in self.warnings:
            lines.append(f"⚠️ {warning}")
        for sample in self.samples:
            lines.append(f"Пример: {sample}")
        return "\n".join(lines)


# ---------------------------------------------------------------- readers (format -> rows)


@dataclass
class Table:
    """A rectangular block of cells plus where it came from."""

    rows: list[list]
    sheet: str = ""
    fmt: str = ""
    source: str = ""


def _zip_kind(data: bytes) -> str:
    """'xlsx' for an Office package, 'zip' for an ordinary archive, '' if unreadable."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return ""
    return "xlsx" if "[Content_Types].xml" in names and any(n.startswith("xl/") for n in names) else "zip"


def read_xlsx(data: bytes, source: str = "") -> list[Table]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables: list[Table] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            if rows:
                tables.append(Table(rows=rows, sheet=worksheet.title, fmt="xlsx", source=source))
    finally:
        workbook.close()
    return tables


def read_csv_text(text: str, source: str = "") -> list[Table]:
    lines = [line for line in text.lstrip("﻿").splitlines() if not line.lstrip().startswith("#")]
    body = "\n".join(lines)
    if not body.strip():
        return []
    head = next((line for line in lines if line.strip()), "")
    delimiter = ";"
    counts = {candidate: head.count(candidate) for candidate in (";", "\t", ",", "|")}
    best = max(counts, key=lambda key: counts[key])
    if counts[best]:
        delimiter = best
    rows = [row for row in csv.reader(io.StringIO(body), delimiter=delimiter)]
    return [Table(rows=rows, fmt="csv", source=source)] if rows else []


def read_dbf(data: bytes, source: str = "") -> list[Table]:
    """Minimal DBF reader - sметные базы are still shipped as dBase files."""
    if len(data) < 32:
        return []
    header_len, record_len = struct.unpack("<HH", data[8:12])
    fields: list[tuple[str, int]] = []
    offset = 32
    while offset < header_len - 1 and data[offset] not in (0x0D, 0x00):
        raw = data[offset:offset + 32]
        name = raw[:11].split(b"\x00")[0]
        fields.append((_decode(name), raw[16]))
        offset += 32
    if not fields:
        return []
    rows: list[list] = [[name for name, _ in fields]]
    position = header_len
    while position + record_len <= len(data):
        record = data[position:position + record_len]
        position += record_len
        if not record or record[0:1] == b"*":      # deleted
            continue
        if record[0:1] == b"\x1a":                  # EOF
            break
        cursor = 1
        values: list = []
        for _, size in fields:
            values.append(_decode(record[cursor:cursor + size]).strip())
            cursor += size
        if any(values):
            rows.append(values)
    return [Table(rows=rows, fmt="dbf", source=source)]


def _decode(raw: bytes) -> str:
    for encoding in ("cp866", "cp1251", "utf-8"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", "replace")


def read_xml(data: bytes, source: str = "") -> list[Table]:
    """Flatten an XML export: find the repeating element that looks like a data row."""
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError as error:
        raise ValueError(f"не удалось разобрать XML: {error}") from error

    candidates: dict[str, list[dict[str, str]]] = {}
    for element in root.iter():
        record: dict[str, str] = {}
        for key, value in element.attrib.items():
            record[key] = value
        for child in element:
            if len(child) == 0 and (child.text or "").strip():
                record[_tag(child)] = (child.text or "").strip()
        if len(record) < 2:
            continue
        mapped = match_columns(list(record.keys()))
        if "name" in mapped or "code" in mapped:
            candidates.setdefault(_tag(element), []).append(record)

    if not candidates:
        return []
    tag = max(candidates, key=lambda key: len(candidates[key]))
    records = candidates[tag]
    headers: list[str] = []
    for record in records:
        for key in record:
            if key not in headers:
                headers.append(key)
    rows: list[list] = [headers]
    rows.extend([record.get(header, "") for header in headers] for record in records)
    return [Table(rows=rows, sheet=tag, fmt="xml", source=source)]


def _tag(element) -> str:
    return element.tag.split("}")[-1]


def read_any(data: bytes | str, filename: str = "") -> list[Table]:
    """Dispatch by content first, extension second. ZIP archives are walked recursively."""
    if isinstance(data, str):
        return read_csv_text(data, source=filename)
    name = (filename or "").lower()
    if data[:4] == b"PK\x03\x04":
        kind = _zip_kind(data)
        if kind == "zip":
            return _read_zip(data, filename)
        if kind == "xlsx":
            return read_xlsx(data, source=filename)
    if name.endswith((".xlsx", ".xlsm")):
        return read_xlsx(data, source=filename)
    if data[:1] in (b"\x03", b"\x04", b"\x05", b"\x30", b"\x83", b"\xf5") and name.endswith(".dbf"):
        return read_dbf(data, source=filename)
    stripped = data.lstrip()[:200]
    if stripped.startswith(b"<?xml") or stripped.startswith(b"<"):
        return read_xml(data, source=filename)
    if name.endswith(".dbf"):
        return read_dbf(data, source=filename)
    return read_csv_text(_decode(data), source=filename)


def _read_zip(data: bytes, filename: str) -> list[Table]:
    tables: list[Table] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for info in archive.infolist():
            if info.is_dir() or info.file_size > 200 * 1024 * 1024:
                continue
            if not info.filename.lower().endswith((".xlsx", ".xlsm", ".csv", ".txt", ".xml", ".dbf")):
                continue
            member = archive.read(info)
            inner = f"{filename}/{info.filename}" if filename else info.filename
            try:
                tables.extend(read_any(member, inner))
            except (ValueError, zipfile.BadZipFile):
                continue
    return tables


# ---------------------------------------------------------------- header detection


def find_header(rows: list[list]) -> tuple[int, dict[str, int], int]:
    """Return (header row index, column map, header depth 1..2).

    Header rows are often split in two ('Эксплуатация машин' / 'в т.ч. оплата труда
    машинистов'), so each column is classified from its deepest non-empty cell upwards.
    """
    best: tuple[int, dict[str, int], int] = (-1, {}, 1)
    best_score = 0
    limit = min(len(rows), MAX_HEADER_SCAN)
    for index in range(limit):
        if not any(classify(cell) for cell in rows[index]):
            continue                      # пустая строка или заголовок таблицы - не шапка
        window = rows[index:index + 3]
        width = max((len(row) for row in window), default=0)
        # depth=1 wins ties: a two-row window is only taken when the second row really
        # adds meaning (a sub-header like «в том числе оплата труда машинистов»).
        for depth in (1, 2):
            if index + depth > len(rows):
                continue
            headers = [_header_cell(rows, index, depth, column) for column in range(width)]
            mapped = match_columns(headers)
            score = len(mapped) + (2 if {"name", "code"} <= set(mapped) else 0)
            if score > best_score and ("name" in mapped or "code" in mapped):
                best_score = score
                best = (index, mapped, depth)
    return best


def _header_cell(rows: list[list], index: int, depth: int, column: int) -> str:
    """The header text of one column: deepest meaningful cell of the header rows."""
    cells = [
        rows[index + offset][column]
        for offset in range(depth)
        if column < len(rows[index + offset]) and rows[index + offset][column] not in (None, "")
    ]
    for cell in reversed(cells):
        if normalize_header(cell) in GENERIC_HEADERS:
            continue                      # «всего» под «Эксплуатация машин» - смысл в родителе
        if classify(cell):
            return str(cell)
    return str(cells[-1]) if cells else ""


# ---------------------------------------------------------------- rates


def rates_from_table(table: Table, *, source: str = "") -> tuple[list[Rate], ImportReport]:
    report = ImportReport(kind="rates", source=table.source or source, fmt=table.fmt, sheet=table.sheet)
    header_row, columns, depth = find_header(table.rows)
    if header_row < 0 or "name" not in columns and "code" not in columns:
        report.warnings.append("не нашёл строку заголовка со столбцами «Шифр»/«Наименование»")
        return [], report
    report.header_row = header_row + 1
    report.columns = {
        key: _header_cell(table.rows, header_row, depth, index) or f"столбец {index + 1}"
        for key, index in sorted(columns.items(), key=lambda pair: pair[1])
    }

    def cell(row: list, key: str) -> str:
        index = columns.get(key, -1)
        if index < 0 or index >= len(row):
            return ""
        value = row[index]
        return "" if value is None else str(value).strip()

    out: list[Rate] = []
    seen: set[str] = set()
    section = ""
    for row in table.rows[header_row + depth:]:
        if not any(value not in (None, "") for value in row):
            continue
        report.rows_seen += 1
        name = cell(row, "name")
        code = cell(row, "code")
        numbers = [cell(row, key) for key in ("ot", "em", "mat", "pz")]
        has_money = any(dec(value) for value in numbers)
        if name and not code and not has_money:
            section = name                      # раздел/подраздел таблицы
            report.skipped["строк-разделов"] += 1
            continue
        if not name:
            report.skipped["без наименования"] += 1
            continue
        if not has_money:
            report.skipped["без стоимостных показателей"] += 1
            continue
        ot, em, zpm, mat = (dec(cell(row, key)) for key in ("ot", "em", "zpm", "mat"))
        pz = dec(cell(row, "pz"))
        if not mat and pz:                       # МАТ отдельным столбцом нет - выводим из ПЗ
            mat = max(pz - ot - em, ZERO)
        work_type = cell(row, "work_type").lower() or sbornik_work_type(code) or ""
        key = code or f"{name}|{cell(row, 'unit')}"
        if key in seen:
            report.skipped["дубликатов шифра"] += 1
            continue
        seen.add(key)
        out.append(
            Rate(
                code=code,
                name=name,
                unit=cell(row, "unit") or "шт",
                ot=ot, em=em, zpm=zpm, mat=mat,
                tz=dec(cell(row, "tz")),
                work_type=work_type or "прочие",
                source=cell(row, "source") or source or table.source or "импорт",
                keywords=" ".join(x for x in (cell(row, "keywords"), section) if x),
            )
        )
        report.imported += 1
        if len(report.samples) < 3:
            report.samples.append(f"{code or '—'} {name[:48]} · {out[-1].unit} · ПЗ {out[-1].direct}")

    if out and not any(rate.work_type != "прочие" for rate in out):
        report.warnings.append(
            "вид работ не определён ни у одной расценки — НР и СП будут нулевыми, "
            "задайте столбец «Вид работ» или дополните data/rates/sbornik_work_type.csv"
        )
    if out and not any(rate.tz for rate in out):
        report.warnings.append("нет затрат труда (чел.-ч) — в смете не будет справки о трудозатратах")
    return out, report


def import_rates(data: bytes | str, filename: str = "") -> tuple[list[Rate], ImportReport]:
    tables = read_any(data, filename)
    if not tables:
        return [], ImportReport(kind="rates", source=filename, warnings=["файл пустой или формат не поддерживается"])
    total = ImportReport(kind="rates", source=filename)
    rates: list[Rate] = []
    codes: set[str] = set()
    for table in tables:
        chunk, report = rates_from_table(table, source=filename)
        total.merge(report)
        for rate in chunk:
            key = rate.code or f"{rate.name}|{rate.unit}"
            if key in codes:
                total.skipped["дубликатов шифра"] += 1
                continue
            codes.add(key)
            rates.append(rate)
    total.imported = len(rates)
    if not rates and not total.warnings:
        total.warnings.append("ни одной расценки не распознано")
    return rates, total


def rates_to_csv(rates: list[Rate]) -> str:
    """Canonical CSV - what the bot stores and the catalog reads back."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    writer.writerow(["Шифр", "Наименование", "Ед.изм.", "ОТ", "ЭМ", "ЗПМ", "МАТ", "ТЗ",
                     "Вид работ", "Источник", "Ключевые слова"])
    for rate in rates:
        writer.writerow([
            rate.code, rate.name, rate.unit, rate.ot, rate.em, rate.zpm, rate.mat, rate.tz,
            rate.work_type, rate.source, rate.keywords,
        ])
    return buffer.getvalue()


# ---------------------------------------------------------------- indices


def indices_from_table(table: Table, *, source: str = "") -> tuple[dict[str, Indices], ImportReport]:
    report = ImportReport(kind="indices", source=table.source or source, fmt=table.fmt, sheet=table.sheet)
    header_row, columns, depth = find_header(table.rows)
    if header_row < 0:
        header_row, columns, depth = _find_index_header(table.rows)
    if header_row < 0:
        report.warnings.append("не нашёл столбцы с индексами (ИОТ/ИЭМ/ИМАТ или «Индекс СМР»)")
        return {}, report
    report.header_row = header_row + 1
    report.columns = {
        key: _header_cell(table.rows, header_row, depth, index) or f"столбец {index + 1}"
        for key, index in sorted(columns.items(), key=lambda pair: pair[1])
    }

    def cell(row: list, key: str) -> str:
        index = columns.get(key, -1)
        if index < 0 or index >= len(row):
            return ""
        value = row[index]
        return "" if value is None else str(value).strip()

    out: dict[str, Indices] = {}
    for row in table.rows[header_row + depth:]:
        if not any(value not in (None, "") for value in row):
            continue
        report.rows_seen += 1
        region = cell(row, "region") or cell(row, "name")
        group = cell(row, "group")
        period = cell(row, "period") or table.sheet
        smr = dec(cell(row, "smr"))
        ot, em, mat = (dec(cell(row, key)) for key in ("ot", "em", "mat"))
        if not any((ot, em, mat, smr)):
            report.skipped["без значений индексов"] += 1
            continue
        if not region:
            report.skipped["без региона"] += 1
            continue
        if not any((ot, em, mat)):
            ot = em = mat = smr                 # единый индекс к СМР
        label = " · ".join(x for x in (region, group, period) if x)
        out[label] = Indices(
            region=region,
            period=period,
            ot=ot or ONE,
            em=em or ONE,
            mat=mat or ONE,
            zpm=dec(cell(row, "zpm"), None) if cell(row, "zpm") else None,
        )
        report.imported += 1
        if len(report.samples) < 3:
            report.samples.append(f"{label}: ИОТ {ot or smr}, ИЭМ {em or smr}, ИМАТ {mat or smr}")
    for label, value in out.items():
        if max(value.ot, value.em, value.mat) > Decimal(100):
            report.warnings.append(f"«{label}»: индекс больше 100 — проверьте, не процент ли это")
            break
    return out, report


def _find_index_header(rows: list[list]) -> tuple[int, dict[str, int], int]:
    """Index tables sometimes have no 'Наименование' column at all."""
    for index in range(min(len(rows), MAX_HEADER_SCAN)):
        mapped = match_columns([str(cell or "") for cell in rows[index]])
        if {"ot", "em", "mat"} <= set(mapped) or "smr" in mapped:
            if "region" not in mapped and "name" not in mapped:
                mapped["region"] = 0
            return index, mapped, 1
    return -1, {}, 1


def import_indices(data: bytes | str, filename: str = "") -> tuple[dict[str, Indices], ImportReport]:
    tables = read_any(data, filename)
    if not tables:
        return {}, ImportReport(kind="indices", source=filename, warnings=["файл пустой или формат не поддерживается"])
    total = ImportReport(kind="indices", source=filename)
    out: dict[str, Indices] = {}
    for table in tables:
        chunk, report = indices_from_table(table, source=filename)
        total.merge(report)
        out.update(chunk)
    total.imported = len(out)
    if not out and not total.warnings:
        total.warnings.append("ни одного индекса не распознано")
    return out, total


def indices_to_csv(indices: dict[str, Indices]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    writer.writerow(["name", "region", "period", "ot", "em", "mat", "zpm"])
    for label, value in indices.items():
        writer.writerow([label, value.region, value.period, value.ot, value.em, value.mat,
                         value.zpm if value.zpm is not None else ""])
    return buffer.getvalue()


# ---------------------------------------------------------------- CLI


def main(argv: list[str] | None = None) -> int:
    import argparse

    parser = argparse.ArgumentParser(
        prog="smeta.importer",
        description="Импорт ФСНБ-2022 (ФЕР/ГЭСН) и индексов ФГИС ЦС из XLSX, XML, DBF, CSV или ZIP",
    )
    parser.add_argument("source", help="файл или архив")
    parser.add_argument("--kind", choices=("rates", "indices", "auto"), default="auto")
    parser.add_argument("--out", help="куда записать канонический CSV")
    parser.add_argument("--limit", type=int, default=0, help="взять только первые N строк (проверка формата)")
    args = parser.parse_args(argv)

    path = Path(args.source)
    data = path.read_bytes()
    kind = args.kind
    if kind == "auto":
        name = path.name.lower()
        kind = "indices" if ("индекс" in name or "index" in name or "ксмр" in name) else "rates"

    if kind == "rates":
        rates, report = import_rates(data, path.name)
        if args.limit:
            rates = rates[: args.limit]
        print(report.as_text())
        if args.out and rates:
            Path(args.out).parent.mkdir(parents=True, exist_ok=True)
            Path(args.out).write_text(rates_to_csv(rates), encoding="utf-8")
            print(f"Записал {len(rates)} расценок в {args.out}")
        return 0 if rates else 1

    indices, report = import_indices(data, path.name)
    print(report.as_text())
    if args.out and indices:
        Path(args.out).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out).write_text(indices_to_csv(indices), encoding="utf-8")
        print(f"Записал {len(indices)} индексов в {args.out}")
    return 0 if indices else 1


if __name__ == "__main__":
    raise SystemExit(main())
