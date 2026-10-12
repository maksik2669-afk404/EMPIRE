"""Проверка чужой сметы: арифметика, повторы, итоги, НДС, цены против справочника, забытые работы.

Бот не придумывает претензий: каждое замечание — это измеренное расхождение или явный
вопрос («есть штукатурка, нет грунтовки»), и у каждого указано, на чём оно основано.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import datafiles
from .catalog import Catalog, Indices, tokens
from .headers import classify, normalize_header
from .importer import read_any
from .model import Position
from .money import ZERO, dec, money, rub
from .units import norm_unit, parse_unit

OVER = Decimal("0.25")       # дороже справочника более чем на 25 % — замечание
UNDER = Decimal("-0.5")      # дешевле более чем вдвое — вопрос к составу работ


def classify_audit(header) -> str | None:
    h = normalize_header(header)
    if not h:
        return None
    if "кол" in h or "объем" in h or "количество" in h:
        return "qty"
    if "цена" in h or "за ед" in h or "стоимость единицы" in h or ("расценка" in h and "шифр" not in h):
        return "price"
    if "сумма" in h or "стоимость" in h or "всего" in h or "итого" in h:
        return "total"
    kind = classify(header)
    return kind if kind in {"name", "code", "unit"} else None


def _find_header(rows: list[list]) -> tuple[int, dict[str, int]]:
    best, best_map = -1, {}
    for index, row in enumerate(rows[:40]):
        mapped: dict[str, int] = {}
        for column, cell in enumerate(row):
            kind = classify_audit(cell)
            if kind and kind not in mapped:
                mapped[kind] = column
        numeric = sum(1 for key in ("qty", "price", "total") if key in mapped)
        if "name" in mapped and numeric >= 2 and len(mapped) > len(best_map):
            best, best_map = index, mapped
    return best, best_map


@dataclass
class AuditRow:
    row: int
    name: str
    unit: str
    qty: Decimal
    price: Decimal
    total: Decimal
    code: str = ""
    reference: Decimal | None = None
    reference_code: str = ""
    issues: list[str] = field(default_factory=list)

    @property
    def deviation(self) -> Decimal | None:
        if not self.reference or not self.price:
            return None
        return ((self.price - self.reference) / self.reference).quantize(Decimal("0.001"))


@dataclass
class AuditReport:
    source: str = ""
    rows: list[AuditRow] = field(default_factory=list)
    summary_rows: list[tuple[str, Decimal]] = field(default_factory=list)
    totals_issues: list[str] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)
    basis: str = ""
    vat_known: bool = False

    @property
    def computed(self) -> Decimal:
        return money(sum((row.total for row in self.rows), ZERO))

    @property
    def flagged(self) -> list[AuditRow]:
        return [row for row in self.rows if row.issues]

    @property
    def overprice(self) -> Decimal:
        """Сколько денег сверх справочника в строках, помеченных как дорогие."""
        excess = ZERO
        for row in self.rows:
            deviation = row.deviation
            if deviation is not None and deviation > OVER:
                excess += (row.price - row.reference) * row.qty
        return money(excess)

    def as_text(self, limit: int = 15) -> str:
        lines = [
            f"🔍 Проверка сметы «{self.source}»",
            f"Строк с работами и материалами: {len(self.rows)}, сумма по строкам: {rub(self.computed)} ₽",
            f"Замечаний по строкам: {len(self.flagged)}",
        ]
        if self.overprice:
            lines.append(f"💸 Дороже справочника в сумме на {rub(self.overprice)} ₽")
        for issue in self.totals_issues:
            lines.append(f"🧮 {issue}")
        if self.flagged:
            lines.append("")
            for row in self.flagged[:limit]:
                lines.append(f"• стр. {row.row} «{row.name[:60]}»")
                for issue in row.issues:
                    lines.append(f"   — {issue}")
            if len(self.flagged) > limit:
                lines.append(f"… и ещё {len(self.flagged) - limit} — все в файле отчёта")
        if self.missing:
            lines.append("")
            lines.append("❓ Возможно, забыты работы:")
            lines.extend(f"• {item}" for item in self.missing)
        lines.append("")
        lines.append(f"Цены сравнивались с: {self.basis}.")
        if not self.vat_known:
            lines.append("В смете нет строки НДС — цены сравнивались как цены без НДС.")
        if not self.flagged and not self.totals_issues and not self.missing:
            lines.append("✅ Ошибок в арифметике, повторов и забытых работ не нашёл.")
        return "\n".join(lines)


def audit(data: bytes | str, filename: str, catalog: Catalog, indices: Indices) -> AuditReport:
    report = AuditReport(source=filename)
    tables = read_any(data, filename)
    header_row, columns, table = -1, {}, None
    for candidate in tables:
        header_row, columns = _find_header(candidate.rows)
        if header_row >= 0:
            table = candidate
            break
    if table is None:
        raise ValueError("не нашёл таблицу сметы: нужны столбцы «Наименование» и хотя бы два из "
                         "«Количество», «Цена», «Сумма»")

    def cell(row: list, key: str) -> str:
        index = columns.get(key, -1)
        if index < 0 or index >= len(row) or row[index] is None:
            return ""
        return str(row[index]).strip()

    seen: dict[tuple, int] = {}
    for offset, row in enumerate(table.rows[header_row + 1:], start=header_row + 2):
        name = cell(row, "name")
        if not name:
            continue
        lowered = name.lower()
        qty, price, total = dec(cell(row, "qty")), dec(cell(row, "price")), dec(cell(row, "total"))
        if lowered.startswith(("итого", "всего", "ндс", "в том числе", "в т.ч")) or "ндс" in lowered.split():
            value = total or price or qty
            if value:
                report.summary_rows.append((name, money(value)))
            continue
        if not any((qty, price, total)):
            continue
        if not total and qty and price:
            total = money(qty * price)
        if not price and qty and total:
            price = money(total / qty)
        item = AuditRow(row=offset, name=name, unit=norm_unit(cell(row, "unit")) or "",
                        qty=qty, price=price, total=money(total), code=cell(row, "code"))
        if qty and price and total:
            expected = money(qty * price)
            if abs(expected - item.total) > max(Decimal(1), item.total * Decimal("0.005")):
                item.issues.append(f"арифметика: {qty} × {rub(price)} = {rub(expected)}, а в смете {rub(item.total)}")
        key = (" ".join(tokens(name)), item.unit, qty)
        if key in seen:
            item.issues.append(f"повтор строки {seen[key]} — возможен двойной счёт")
        else:
            seen[key] = offset
        _compare_with_catalog(item, catalog, indices)
        report.rows.append(item)

    _check_totals(report)
    report.missing = _missing_works(report.rows)
    region = indices.region or "текущие цены"
    report.basis = f"ваш справочник расценок ({len(catalog.rates)} шт.), уровень цен «{region}», с НР и СП"
    return report


def _compare_with_catalog(item: AuditRow, catalog: Catalog, indices: Indices) -> None:
    rate = catalog.by_code(item.code) if item.code else None
    if rate is None:
        found = catalog.best(item.name, threshold=0.6)
        rate = found[0] if found else None
    if rate is None or not item.price:
        return
    if item.unit and norm_unit(item.unit) != parse_unit(rate.unit)[1]:
        return                                   # другая единица — сравнивать нельзя
    position = Position.from_rate(rate, 1, catalog.norms(rate.work_type))
    reference = position.total(indices)
    if reference <= ZERO:
        return
    item.reference, item.reference_code = reference, rate.code
    deviation = item.deviation
    if deviation is not None and deviation > OVER:
        item.issues.append(f"дороже справочника на {deviation * 100:.0f}%: {rub(item.price)} против "
                           f"{rub(reference)} ₽/{item.unit or 'ед.'} ({rate.code})")
    elif deviation is not None and deviation < UNDER:
        item.issues.append(f"дешевле справочника на {-deviation * 100:.0f}% — уточните состав работ ({rate.code})")


def _check_totals(report: AuditReport) -> None:
    computed = report.computed
    declared = next((value for name, value in report.summary_rows if name.lower().startswith("итого")), None)
    if declared is not None and abs(declared - computed) > Decimal(1):
        report.totals_issues.append(f"«Итого» в смете {rub(declared)} ₽, а сумма строк {rub(computed)} ₽ "
                                    f"(разница {rub(declared - computed)} ₽)")
    vat = next((value for name, value in report.summary_rows if "ндс" in name.lower()), None)
    if vat is not None:
        report.vat_known = True
        base = declared if declared is not None else computed
        expected = money(base * Decimal("0.2"))
        if abs(vat - expected) > Decimal(1) and abs(vat - money(base * Decimal("0.2") / Decimal("1.2"))) > Decimal(1):
            report.totals_issues.append(f"НДС {rub(vat)} ₽ не равен 20 % от итога ({rub(expected)} ₽)")


def _missing_works(rows: list[AuditRow]) -> list[str]:
    present = [set(tokens(row.name)) for row in rows]
    out: list[str] = []
    for rule in datafiles.companions():
        if_has = [t for word in rule.get("if_has", "").split() for t in tokens(word)]
        expect = [t for word in rule.get("expect", "").split() for t in tokens(word)]
        has_first = any(any(t in words for t in if_has) for words in present)
        has_second = any(any(t in words for t in expect) for words in present)
        if has_first and not has_second:
            out.append(f"нет «{rule['expect'].split()[0]}», хотя есть «{rule['if_has'].split()[0]}» — {rule.get('why', '')}")
    return out


def report_xlsx(report: AuditReport, path: str | Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Проверка"
    sheet.append([f"Проверка сметы «{report.source}»"])
    sheet["A1"].font = Font(size=12, bold=True)
    sheet.append([f"Сравнение цен: {report.basis}"])
    for issue in report.totals_issues:
        sheet.append([issue])
    sheet.append([])
    headers = ["Стр.", "Наименование", "Ед.", "Кол-во", "Цена", "Сумма", "Справочник, ₽/ед.",
               "Отклонение", "Шифр справочника", "Замечания"]
    sheet.append(headers)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    red = PatternFill("solid", fgColor="FFFFE5E5")
    for row in report.rows:
        deviation = row.deviation
        sheet.append([row.row, row.name, row.unit, float(row.qty), float(row.price), float(row.total),
                      float(row.reference) if row.reference else None,
                      f"{deviation * 100:+.0f}%" if deviation is not None else "",
                      row.reference_code, "; ".join(row.issues)])
        if row.issues:
            for cell in sheet[sheet.max_row]:
                cell.fill = red
        sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(wrap_text=True)
        sheet.cell(row=sheet.max_row, column=10).alignment = Alignment(wrap_text=True)
    if report.missing:
        sheet.append([])
        sheet.append(["Возможно, забыты работы:"])
        for item in report.missing:
            sheet.append(["", item])
    for column, width in zip("ABCDEFGHIJ", (6, 50, 8, 10, 12, 14, 14, 11, 16, 60)):
        sheet.column_dimensions[column].width = width
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return out
