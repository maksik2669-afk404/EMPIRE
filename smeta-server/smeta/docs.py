"""Документы из сметы: КС-2, КС-3, наряд-задание, допсоглашение, КП, пересчёт.

Унифицированные формы КС-2/КС-3 с 2013 года не обязательны (402-ФЗ «О бухучёте»), но
заказчики привыкли к их составу граф — его документы и повторяют.
"""
from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import KIND_TITLES, MATERIAL, WORK, Estimate
from .money import ZERO, dec, money, pct, rub
from .progress import SCHEDULE_PRESETS, year_to_date
from .reprice import RepriceResult
from .versions import Diff

THIN = Side(style="thin", color="FF808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="FFEFEFEF")
MONEY = "# ##0.00"


class Sheet:
    """Tiny helper for form-like sheets: titles, labelled lines, bordered tables."""

    def __init__(self, title: str, widths: list[int]):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = title
        self.width = len(widths)
        self.row = 1
        for index, value in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(index)].width = value
        self.ws.page_setup.orientation = "landscape"
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.page_setup.fitToWidth = 1

    def text(self, value: str, *, bold=False, size=10, align="left", italic=False):
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.width)
        cell = self.ws.cell(row=self.row, column=1, value=value)
        cell.font = Font(size=size, bold=bold, italic=italic)
        cell.alignment = Alignment(horizontal=align, wrap_text=True, vertical="center")
        self.row += 1

    def blank(self, count: int = 1):
        self.row += count

    def table(self, header: list[str], rows: list[list], money_cols: set[int] = frozenset(), qty_cols: set[int] = frozenset(),
              wrap_cols: set[int] = frozenset()):
        for column, value in enumerate(header, 1):
            cell = self.ws.cell(row=self.row, column=column, value=value)
            cell.font = Font(size=9, bold=True)
            cell.fill = HEAD
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.row += 1
        for values in rows:
            for column, value in enumerate(values, 1):
                if isinstance(value, Decimal):
                    value = float(value)
                cell = self.ws.cell(row=self.row, column=column, value=value)
                cell.font = Font(size=9)
                cell.border = BOX
                cell.alignment = Alignment(vertical="top", wrap_text=column in wrap_cols)
                if column in money_cols:
                    cell.number_format = MONEY
                if column in qty_cols:
                    cell.number_format = "# ##0.###"
            self.row += 1

    def total(self, label: str, value, *, bold=False, value_col: int | None = None):
        value_col = value_col or self.width
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=value_col - 1)
        cell = self.ws.cell(row=self.row, column=1, value=label)
        cell.font = Font(size=10, bold=bold)
        cell.alignment = Alignment(horizontal="right")
        target = self.ws.cell(row=self.row, column=value_col, value=float(money(value)))
        target.font = Font(size=10, bold=bold)
        target.number_format = MONEY
        target.border = BOX
        self.row += 1

    def signatures(self, left: str, right: str):
        self.blank()
        half = max(self.width // 2, 1)
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=half)
        self.ws.cell(row=self.row, column=1, value=left).font = Font(size=10)
        self.ws.merge_cells(start_row=self.row, start_column=half + 1, end_row=self.row, end_column=self.width)
        self.ws.cell(row=self.row, column=half + 1, value=right).font = Font(size=10)
        self.row += 1

    def save(self, path: str | Path) -> Path:
        out = Path(path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(out)
        return out


def _parties(sheet: Sheet, estimate: Estimate) -> None:
    company = estimate.company
    contractor = ", ".join(x for x in (company.name, f"ИНН {company.inn}" if company.inn else "", company.address) if x)
    for label, value in (("Заказчик", estimate.customer), ("Подрядчик", contractor),
                         ("Объект", estimate.object_name), ("Договор", estimate.contract)):
        sheet.text(f"{label}: {value or '—'}")


def _d(value: str) -> str:
    return date.fromisoformat(value).strftime("%d.%m.%Y") if value else "—"


# ---------------------------------------------------------------- КС-2


def ks2_xlsx(estimate: Estimate, act: dict, path: str | Path) -> Path:
    sheet = Sheet("КС-2", [6, 9, 52, 18, 9, 11, 14, 16])
    sheet.text("Форма КС-2", align="right", size=8, italic=True)
    _parties(sheet, estimate)
    sheet.blank()
    sheet.text(f"АКТ № {act['number']} о приёмке выполненных работ", bold=True, size=13, align="center")
    sheet.text(f"Дата составления: {_d(act['date'])}    Отчётный период: с {_d(act['from'])} по {_d(act['to'])}",
               align="center")
    sheet.text(f"Сметная стоимость по смете № {estimate.number or '—'}: {rub(estimate.total())} ₽", align="center")
    sheet.blank()
    rows = []
    for number, (index, volume, amount) in enumerate(act["items"], 1):
        position = estimate.positions[index - 1]
        volume, amount = dec(volume), dec(amount)
        rows.append([number, index, position.name, position.code or "договорная цена", position.unit,
                     volume, money(amount / volume) if volume else ZERO, amount])
    sheet.table(["№ п/п", "№ поз. по смете", "Наименование работ", "Обоснование", "Ед. изм.",
                 "Количество", "Цена за ед., ₽", "Стоимость, ₽"], rows,
                money_cols={7, 8}, qty_cols={6}, wrap_cols={3})
    sheet.blank()
    sheet.total("Итого по акту", dec(act["smr"]), bold=True)
    if dec(act["limited"]):
        sheet.total("Лимитированные затраты", dec(act["limited"]))
    if estimate.vat_pct:
        sheet.total(f"НДС {estimate.vat_pct:g}%", dec(act["vat"]))
    sheet.total("Всего по акту", dec(act["total"]), bold=True)
    sheet.signatures(f"Сдал: {estimate.company.signer_title} ____________ {estimate.company.signer}",
                     "Принял: ____________")
    return sheet.save(path)


# ---------------------------------------------------------------- КС-3


def ks3_xlsx(estimate: Estimate, act: dict, path: str | Path) -> Path:
    in_year, since_start = year_to_date(estimate, act)
    rate = Decimal(1) + pct(estimate.vat_pct)
    sheet = Sheet("КС-3", [6, 60, 16, 20, 20, 22])
    sheet.text("Форма КС-3", align="right", size=8, italic=True)
    _parties(sheet, estimate)
    sheet.blank()
    sheet.text(f"СПРАВКА № {act['number']} о стоимости выполненных работ и затрат", bold=True, size=13, align="center")
    sheet.text(f"Дата составления: {_d(act['date'])}    Отчётный период: с {_d(act['from'])} по {_d(act['to'])}",
               align="center")
    sheet.blank()
    this_net = money(dec(act["smr"]) + dec(act["limited"]))
    rows = [
        [1, "Всего работ и затрат, включаемых в стоимость работ (без НДС)", "",
         money(since_start / rate), money(in_year / rate), this_net],
        [2, f"в том числе по смете № {estimate.number or '—'}", estimate.number,
         money(since_start / rate), money(in_year / rate), this_net],
    ]
    sheet.table(["№", "Наименование работ и затрат", "Код",
                 "С начала проведения работ, ₽", "С начала года, ₽", "В том числе за отчётный период, ₽"],
                rows, money_cols={4, 5, 6}, wrap_cols={2})
    sheet.blank()
    sheet.total("Итого за отчётный период", this_net, bold=True)
    if estimate.vat_pct:
        sheet.total(f"НДС {estimate.vat_pct:g}%", dec(act["vat"]))
    sheet.total("Всего с учётом НДС", dec(act["total"]), bold=True)
    sheet.signatures(f"Подрядчик: {estimate.company.signer_title} ____________ {estimate.company.signer}",
                     "Заказчик: ____________")
    return sheet.save(path)


# ---------------------------------------------------------------- наряд


def crew_rows(estimate: Estimate) -> tuple[list[list], Decimal, int]:
    rows, total, missing = [], ZERO, 0
    number = 0
    for position in estimate.positions:
        if position.kind != WORK:
            continue
        number += 1
        if position.crew_price is None:
            missing += 1
            rows.append([number, position.name, position.unit, position.qty, None, None, ""])
            continue
        amount = money(position.crew_price * position.qty)
        total += amount
        rows.append([number, position.name, position.unit, position.qty, position.crew_price, amount, ""])
    return rows, money(total), missing


def naryad_xlsx(estimate: Estimate, path: str | Path) -> Path:
    rows, total, missing = crew_rows(estimate)
    sheet = Sheet("Наряд", [6, 56, 9, 11, 16, 16, 20])
    sheet.text(f"НАРЯД-ЗАДАНИЕ № {estimate.number or '—'}-Н", bold=True, size=13, align="center")
    sheet.text(f"Объект: {estimate.object_name or '—'}")
    sheet.text(f"Бригада: {estimate.crew or '____________'}    Дата выдачи: {date.today().strftime('%d.%m.%Y')}")
    sheet.blank()
    sheet.table(["№", "Работа", "Ед.", "Объём", "Ставка, ₽/ед.", "Сумма, ₽", "Выполнено (подпись)"],
                rows, money_cols={5, 6}, qty_cols={4}, wrap_cols={2})
    sheet.blank()
    sheet.total("Итого к выплате бригаде", total, bold=True, value_col=6)
    if missing:
        sheet.text(f"Без ставки: {missing} поз. — сумма по ним не начислена", italic=True)
    sheet.signatures(f"Выдал: {estimate.company.signer_title} ____________ {estimate.company.signer}",
                     "Принял бригадир: ____________")
    return sheet.save(path)


# ---------------------------------------------------------------- допсоглашение


def ds_xlsx(estimate: Estimate, change: Diff, number: int, path: str | Path) -> Path:
    sheet = Sheet("ДС", [6, 60, 9, 12, 18])
    sheet.text(f"ДОПОЛНИТЕЛЬНОЕ СОГЛАШЕНИЕ № {number}", bold=True, size=13, align="center")
    sheet.text(f"к договору {estimate.contract or '№ ____ от ____'}", align="center")
    sheet.text(f"г. ____________                                         {date.today().strftime('%d.%m.%Y')}")
    sheet.blank()
    sheet.text(f"{estimate.company.name or 'Подрядчик'} (Подрядчик) и {estimate.customer or 'Заказчик'} (Заказчик) "
               "договорились о следующем:")
    sheet.blank()
    if change.extra:
        sheet.text("1. Подрядчик выполняет дополнительные работы:", bold=True)
        sheet.table(["№", "Наименование работ", "Ед.", "Количество", "Стоимость, ₽"],
                    [[n, c.name, c.unit, c.delta_qty if c.old_qty else c.new_qty, c.delta]
                     for n, c in enumerate(change.extra, 1)], money_cols={5}, qty_cols={4}, wrap_cols={2})
        sheet.total("Итого дополнительных работ", sum((c.delta for c in change.extra), ZERO), value_col=5)
        sheet.blank()
    if change.excluded:
        sheet.text(f"{2 if change.extra else 1}. Из состава работ исключаются:", bold=True)
        sheet.table(["№", "Наименование работ", "Ед.", "Количество", "Уменьшение, ₽"],
                    [[n, c.name, c.unit, -c.delta_qty if c.new_qty else c.old_qty, -c.delta]
                     for n, c in enumerate(change.excluded, 1)], money_cols={5}, qty_cols={4}, wrap_cols={2})
        sheet.blank()
    sign = "увеличивается" if change.delta >= 0 else "уменьшается"
    sheet.text(f"Цена договора {sign} на {rub(abs(change.delta))} ₽"
               f"{' с НДС' if estimate.vat_pct else ''} и составляет {rub(change.new_total)} ₽.")
    sheet.text("Остальные условия договора остаются без изменений.")
    sheet.signatures(f"Подрядчик: ____________ {estimate.company.signer}", "Заказчик: ____________")
    return sheet.save(path)


# ---------------------------------------------------------------- пересчёт


def reprice_xlsx(result: RepriceResult, path: str | Path) -> Path:
    sheet = Sheet("Пересчёт", [6, 60, 18, 18, 16, 10])
    sheet.text(f"Пересчёт сметы: {result.reason}", bold=True, size=12)
    sheet.blank()
    sheet.table(["№", "Позиция", "Было, ₽", "Стало, ₽", "Разница, ₽", "%"],
                [[n, row.name, row.old, row.new, row.delta, f"{row.pct:+}%"] for n, row in enumerate(result.rows, 1)],
                money_cols={3, 4, 5}, wrap_cols={2})
    sheet.blank()
    sheet.total("Было всего", result.old_total, value_col=4)
    sheet.total("Стало всего", result.new_total, bold=True, value_col=4)
    sheet.total("Разница", result.delta, value_col=4)
    return sheet.save(path)


# ---------------------------------------------------------------- КП (PDF)


def kp_groups(estimate: Estimate) -> list[tuple[str, Decimal]]:
    """Amounts by section (or by kind) with limited costs and VAT spread proportionally."""
    scale = estimate.scale()
    groups: dict[str, Decimal] = {}
    use_sections = any(position.section for position in estimate.positions)
    for position in estimate.positions:
        key = (position.section or "Прочие работы") if use_sections else KIND_TITLES.get(position.kind, "Работы")
        groups[key] = groups.get(key, ZERO) + position.total(estimate.indices)
    rows = [(name, money(value * scale)) for name, value in groups.items()]
    drift = estimate.total() - sum((value for _, value in rows), ZERO)
    if rows and drift:
        name, value = rows[-1]
        rows[-1] = (name, money(value + drift))           # копейки округления — в последнюю строку
    return rows


def kp_pdf(estimate: Estimate, path: str | Path, logo: Path | None = None) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .export_pdf import _register_fonts

    font, bold = _register_fonts()
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm,
                            title=f"КП {estimate.number}", author=estimate.company.name)
    body = ParagraphStyle("b", fontName=font, fontSize=10, leading=14)
    small = ParagraphStyle("s", fontName=font, fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    title = ParagraphStyle("t", fontName=bold, fontSize=16, leading=20, spaceAfter=6)
    right = ParagraphStyle("r", parent=body, alignment=2)
    story: list = []

    header_text = "<br/>".join(estimate.company.header_lines()) or "&nbsp;"
    if logo and Path(logo).exists():
        image = Image(str(logo))
        ratio = image.imageHeight / image.imageWidth if image.imageWidth else 1
        image.drawWidth, image.drawHeight = 40 * mm, 40 * mm * ratio
        if image.drawHeight > 25 * mm:
            image.drawHeight, image.drawWidth = 25 * mm, 25 * mm / ratio
        head = Table([[image, Paragraph(header_text, small)]], colWidths=[50 * mm, 124 * mm])
        head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story.append(head)
    else:
        story.append(Paragraph(header_text, small))
    story.append(Spacer(1, 10 * mm))
    story.append(Paragraph("Коммерческое предложение", title))
    story.append(Paragraph(f"№ {estimate.number or '—'} от {estimate.date.strftime('%d.%m.%Y')}", small))
    story.append(Spacer(1, 5 * mm))
    for label, value in (("Заказчик", estimate.customer), ("Объект", estimate.object_name)):
        if value:
            story.append(Paragraph(f"<b>{label}:</b> {value}", body))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Предлагаем выполнить работы на следующих условиях:", body))
    story.append(Spacer(1, 3 * mm))

    rows = [[Paragraph("<b>Состав работ</b>", body), Paragraph("<b>Стоимость, ₽</b>", right)]]
    for name, value in kp_groups(estimate):
        rows.append([Paragraph(name, body), Paragraph(rub(value), right)])
    vat_note = f"в т.ч. НДС {estimate.vat_pct:g}% — {rub(estimate.totals().vat)} ₽" if estimate.vat_pct else "НДС не облагается"
    rows.append([Paragraph(f"<b>Итого</b> ({vat_note})", body), Paragraph(f"<b>{rub(estimate.total())}</b>", right)])
    table = Table(rows, colWidths=[124 * mm, 50 * mm])
    table.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#333333")),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.HexColor("#333333")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    story.append(Spacer(1, 6 * mm))

    kp = estimate.kp or {}
    valid_days = int(dec(kp.get("valid_days") or 14))
    valid_to = (estimate.date + timedelta(days=valid_days)).strftime("%d.%m.%Y")
    if estimate.schedule:
        payment = ", ".join(f"{item['name'].lower()} {item['pct']}%" for item in estimate.schedule)
    else:
        payment = kp.get("payment") or "аванс 30 %, остаток после подписания акта"
    for label, value in (
        ("Срок выполнения", kp.get("term") or "по согласованию, от даты аванса"),
        ("Условия оплаты", payment),
        ("Гарантия на работы", kp.get("warranty") or "12 месяцев"),
        ("Предложение действительно", f"{valid_days} дней, до {valid_to}"),
    ):
        story.append(Paragraph(f"<b>{label}:</b> {value}", body))
    story.append(Spacer(1, 12 * mm))
    signer = f"{estimate.company.signer_title} {estimate.company.name}".strip()
    story.append(Paragraph(f"{signer} ____________ {estimate.company.signer}", body))
    doc.build(story)
    return out


def schedule_presets() -> list[str]:
    return list(SCHEDULE_PRESETS)


__all__ = ["ks2_xlsx", "ks3_xlsx", "naryad_xlsx", "ds_xlsx", "reprice_xlsx", "kp_pdf", "kp_groups", "crew_rows",
           "MATERIAL"]
