"""Import of official data: FSNB-2022 rates and FGIS CS indices from XLSX, XML, DBF, ZIP, CSV."""
import struct
import zipfile
from decimal import Decimal

import pytest
from openpyxl import Workbook

from smeta import fgis, importer
from smeta.catalog import Catalog
from smeta.headers import classify, normalize_header, sbornik_work_type


# ---------- fixtures ----------


def fsnb_xlsx(path):
    """A fragment shaped like a published FER collection: title, two-row header, sections."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ФЕР 15"
    sheet.append(["Федеральные единичные расценки. Сборник 15. Отделочные работы"])
    sheet.append([])
    sheet.append(["Шифр расценки", "Наименование и характеристика работ", "Ед. изм.",
                  "Прямые затраты, руб.", "Оплата труда рабочих, руб.", "Эксплуатация машин, руб.",
                  None, "Материалы, руб.", "Затраты труда рабочих, чел.-ч"])
    sheet.append([None, None, None, None, None, "всего", "в том числе оплата труда машинистов", None, None])
    sheet.append([None, "Раздел 1. Штукатурные работы"])
    sheet.append(["ФЕР15-02-016-03", "Штукатурка поверхностей внутри здания стен", "100 м2",
                  48230.11, 44120.50, 2500.00, 700.00, 1609.61, 110.4])
    sheet.append(["ФЕР15-04-005-01", "Окраска водоэмульсионными составами стен", "100 м2",
                  26100.00, 25000.00, 0, 0, 1100.00, 62.1])
    sheet.append([None, "Раздел 2. Облицовочные работы"])
    sheet.append(["ФЕР15-01-019-05", "Облицовка стен керамическими плитками", "100 м2",
                  152300.44, 148000.00, 1200.00, 300.00, 3100.44, 360.5])
    sheet.append([None, "итого по сборнику", None, None, None, None, None, None, None])
    workbook.save(path)
    return path


FER_XML = """<?xml version="1.0" encoding="utf-8"?>
<База Версия="ФСНБ-2022">
  <Расценка Шифр="ФЕР08-02-001-01" Ед_изм="100 м2">
    <Наименование>Кладка перегородок из кирпича</Наименование>
    <ОплатаТрудаРабочих>31000.40</ОплатаТрудаРабочих>
    <ЭксплуатацияМашин>2100.00</ЭксплуатацияМашин>
    <ОплатаТрудаМашинистов>640.00</ОплатаТрудаМашинистов>
    <Материалы>18400.10</Материалы>
    <ЗатратыТрудаРабочих>310.2</ЗатратыТрудаРабочих>
  </Расценка>
  <Расценка Шифр="ФЕР11-01-011-01" Ед_изм="100 м2">
    <Наименование>Устройство стяжек цементных</Наименование>
    <ОплатаТрудаРабочих>12000.00</ОплатаТрудаРабочих>
    <ЭксплуатацияМашин>3400.00</ЭксплуатацияМашин>
    <ОплатаТрудаМашинистов>1020.00</ОплатаТрудаМашинистов>
    <Материалы>9100.00</Материалы>
    <ЗатратыТрудаРабочих>95.0</ЗатратыТрудаРабочих>
  </Расценка>
</База>"""


def fer_dbf() -> bytes:
    fields = [("SHIFR", 20), ("NAIM", 40), ("EDIZM", 10), ("OT", 12), ("EM", 12),
              ("ZPM", 12), ("MAT", 12), ("TZ", 10)]
    header_len = 32 + 32 * len(fields) + 1
    record_len = 1 + sum(size for _, size in fields)
    out = bytearray(struct.pack("<BBBBIHH20x", 3, 126, 1, 1, 2, header_len, record_len))
    for name, size in fields:
        out += name.encode("cp866").ljust(11, b"\x00") + b"C" + b"\x00" * 4 + bytes([size]) + b"\x00" * 15
    out += b"\x0d"
    rows = [
        ("ФЕР06-01-001-01", "Устройство бетонной подготовки", "100 м3", "23000.00", "8900.00",
         "2670.00", "41000.00", "220.5"),
        ("ФЕР12-01-015-02", "Устройство кровель из металлочерепицы", "100 м2", "35000.00", "2100.00",
         "630.00", "8000.00", "150.0"),
    ]
    for row in rows:
        out += b" "
        for (_, size), value in zip(fields, row):
            out += value.encode("cp866")[:size].ljust(size, b" ")
    out += b"\x1a"
    return bytes(out)


def indices_xlsx(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "IV кв 2026"
    sheet.append(["Индексы изменения сметной стоимости по элементам затрат, IV квартал 2026 г."])
    sheet.append(["Субъект Российской Федерации", "Вид строительства", "ИОТ", "ИЭМ", "ИМАТ"])
    sheet.append(["г. Москва", "Многоквартирные жилые дома", 1.42, 1.33, 1.28])
    sheet.append(["Московская область", "Многоквартирные жилые дома", 1.38, 1.30, 1.25])
    workbook.save(path)
    return path


# ---------- headers ----------


def test_normalize_splits_camel_case_and_underscores():
    assert normalize_header("ОплатаТрудаРабочих") == "оплата труда рабочих"
    assert normalize_header("Ед_изм.") == "ед изм"
    assert normalize_header("Затраты труда\nрабочих, чел.-ч") == "затраты труда рабочих чел -ч"


@pytest.mark.parametrize(
    "header,expected",
    [
        ("Шифр расценки", "code"),
        ("Обоснование", "code"),
        ("Наименование и характеристика работ", "name"),
        ("Ед. изм.", "unit"),
        ("Прямые затраты, руб.", "pz"),
        ("Оплата труда рабочих, руб.", "ot"),
        ("Эксплуатация машин, в т.ч. оплата труда машинистов", "em"),
        ("в том числе оплата труда машинистов", "zpm"),
        ("Материалы, руб.", "mat"),
        ("Затраты труда рабочих, чел.-ч", "tz"),
        ("Затраты труда машинистов, чел.-ч", "tzm"),
        ("Субъект Российской Федерации", "region"),
        ("ИОТ", "ot"),
        ("Индекс СМР", "smr"),
        ("SHIFR", "code"),
        ("NAIM", "name"),
        ("всего", "pz"),
        ("Примечание", None),
    ],
)
def test_classify(header, expected):
    assert classify(header) == expected


def test_sbornik_to_work_type():
    assert sbornik_work_type("ФЕР15-02-016-03") == "отделочные"
    assert sbornik_work_type("ФЕР08-02-001-01") == "каменные"
    assert sbornik_work_type("ГЭСН11-01-011-01") == "полы"
    assert sbornik_work_type("нет цифр") == ""


# ---------- rates ----------


def test_xlsx_two_row_header_and_sections(tmp_path):
    path = fsnb_xlsx(tmp_path / "ФСНБ-2022_сборник15.xlsx")
    rates, report = importer.import_rates(path.read_bytes(), path.name)
    assert len(rates) == 3
    assert report.columns["zpm"] == "в том числе оплата труда машинистов"
    assert report.columns["em"] == "Эксплуатация машин, руб."       # «всего» не перебивает родителя
    assert report.skipped["строк-разделов"] == 3
    first = rates[0]
    assert first.code == "ФЕР15-02-016-03"
    assert first.ot == Decimal("44120.50") and first.zpm == Decimal("700.00")
    assert first.work_type == "отделочные"                          # по номеру сборника
    assert "Раздел 1" in first.keywords                             # раздел попал в поиск


def test_materials_derived_from_direct_costs(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Шифр", "Наименование", "Ед. изм.", "Прямые затраты, руб.",
                  "Оплата труда рабочих", "Эксплуатация машин"])
    sheet.append(["ФЕР01-01-001-01", "Разработка грунта", "1000 м3", 100000, 20000, 30000])
    path = tmp_path / "r.xlsx"
    workbook.save(path)
    rates, _ = importer.import_rates(path.read_bytes(), path.name)
    assert rates[0].mat == Decimal(50000)       # 100000 − 20000 − 30000


def test_xml_import():
    rates, report = importer.import_rates(FER_XML.encode("utf-8"), "fer.xml")
    assert [rate.code for rate in rates] == ["ФЕР08-02-001-01", "ФЕР11-01-011-01"]
    assert rates[0].ot == Decimal("31000.40")
    assert rates[0].unit == "100 м2"
    assert report.fmt == "xml"


def test_dbf_import():
    rates, report = importer.import_rates(fer_dbf(), "fer.dbf")
    assert len(rates) == 2
    assert rates[0].name == "Устройство бетонной подготовки"
    assert rates[1].work_type == "кровельные"
    assert report.fmt == "dbf"


def test_zip_merges_parts(tmp_path):
    xlsx = fsnb_xlsx(tmp_path / "part.xlsx")
    archive = tmp_path / "ФСНБ-2022.zip"
    with zipfile.ZipFile(archive, "w") as zipped:
        zipped.write(xlsx, "part1/сборник15.xlsx")
        zipped.writestr("part2/fer.xml", FER_XML)
    rates, report = importer.import_rates(archive.read_bytes(), archive.name)
    assert len(rates) == 5
    assert report.imported == 5


def test_duplicates_are_reported_once(tmp_path):
    xlsx = fsnb_xlsx(tmp_path / "part.xlsx")
    archive = tmp_path / "twice.zip"
    with zipfile.ZipFile(archive, "w") as zipped:
        zipped.write(xlsx, "a/part.xlsx")
        zipped.write(xlsx, "b/part.xlsx")
    rates, report = importer.import_rates(archive.read_bytes(), archive.name)
    assert len(rates) == 3
    assert report.skipped["дубликатов шифра"] == 3


def test_junk_file_is_rejected_with_a_reason():
    rates, report = importer.import_rates(b"abc;def\n1;2\n", "junk.csv")
    assert not rates
    assert report.warnings


def test_canonical_csv_round_trips_into_the_catalog(tmp_path):
    path = fsnb_xlsx(tmp_path / "ФСНБ.xlsx")
    rates, _ = importer.import_rates(path.read_bytes(), path.name)
    catalog = Catalog.from_csv(importer.rates_to_csv(rates))
    assert len(catalog.rates) == 3
    rate, score = catalog.best("штукатурка стен")
    assert rate.code == "ФЕР15-02-016-03" and score > 0.5
    assert catalog.by_code("ФЕР15-01-019-05").work_type == "отделочные"


# ---------- indices ----------


def test_indices_import(tmp_path):
    path = indices_xlsx(tmp_path / "индексы_IV_кв_2026.xlsx")
    indices, report = importer.import_indices(path.read_bytes(), path.name)
    assert len(indices) == 2
    key = next(iter(indices))
    assert "Москва" in key and "IV кв 2026" in key
    assert indices[key].ot == Decimal("1.42")
    assert report.columns["region"] == "Субъект Российской Федерации"


def test_single_smr_index_fills_all_elements():
    csv_text = "Регион;Индекс СМР\nг. Москва;8.75\n"
    indices, _ = importer.import_indices(csv_text, "ksmr.csv")
    value = next(iter(indices.values()))
    assert value.ot == value.em == value.mat == Decimal("8.75")


def test_percent_looking_indices_raise_a_warning():
    indices, report = importer.import_indices("Регион;ИОТ;ИЭМ;ИМАТ\nг. Москва;142;133;128\n", "i.csv")
    assert indices
    assert any("больше 100" in warning for warning in report.warnings)


# ---------- fgis CLI ----------


def test_fgis_load_writes_canonical_csv(tmp_path):
    path = fsnb_xlsx(tmp_path / "ФСНБ.xlsx")
    out = tmp_path / "fer.csv"
    assert fgis.load(path, "rates", out=out) == 0
    body = out.read_text(encoding="utf-8")
    assert body.startswith("Шифр;Наименование")
    assert "ФЕР15-02-016-03" in body
    assert len(Catalog.from_csv(body).rates) == 3


def test_fgis_load_appends(tmp_path):
    path = fsnb_xlsx(tmp_path / "ФСНБ.xlsx")
    out = tmp_path / "fer.csv"
    fgis.load(path, "rates", out=out)
    (tmp_path / "fer2.xml").write_text(FER_XML, encoding="utf-8")
    assert fgis.load(tmp_path / "fer2.xml", "rates", out=out, append=True) == 0
    assert len(Catalog.from_csv(out.read_text(encoding="utf-8")).rates) == 5


def test_fgis_load_detects_indices_by_name(tmp_path):
    path = indices_xlsx(tmp_path / "индексы.xlsx")
    out = tmp_path / "idx.csv"
    assert fgis.load(path, "auto", out=out) == 0
    assert out.read_text(encoding="utf-8").startswith("name;region;period")


def test_fgis_check_reports_unreachable_host_without_crashing():
    assert fgis.check("https://example.invalid/никуда", timeout=5) == 1
