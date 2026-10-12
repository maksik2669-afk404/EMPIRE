"""Feature modules: progress & acts, obmer, templates, materials by norms, audit, versions, reprice,
coefficients, limited costs, billing, funnel, reverse tax pricing, documents."""
import datetime
import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from smeta import audit, billing, consumption, docs, funnel, obmer, progress, reprice, tax, templates, versions
from smeta.builder import build_from_text
from smeta.catalog import Catalog
from smeta.materials import OfferBook
from smeta.model import MATERIAL, Company, Estimate, Position
from smeta.money import money
from smeta.store import Store


@pytest.fixture(scope="module")
def catalog():
    return Catalog.load()


@pytest.fixture(scope="module")
def offers():
    return OfferBook.load("data/suppliers/demo.csv")


def make(catalog, text="штукатурка стен 120 м2\nстяжка пола 85 м2\nмат: керамогранит 22 м2 1550", **kw):
    estimate = Estimate(number="СМ-1", indices=catalog.index(), company=Company(name="ООО", inn="7707083893"), **kw)
    for match in build_from_text(text, catalog):
        estimate.add(match.position)
    return estimate


# ---------------------------------------------------------------- coefficients and limited costs


def test_coefficients_raise_costs_and_keep_basis(catalog):
    estimate = make(catalog, "штукатурка стен 100 м2")
    before = estimate.total()
    floor_before = estimate.floor()
    estimate.positions[0].apply_coefficient("Стеснённые", "1.15", "1.15", "1", "прил. 10")
    assert estimate.total() > before
    assert estimate.floor() > floor_before                  # коэффициент - реальное удорожание
    assert "ОТ 1.15" in estimate.positions[0].k_basis
    estimate.positions[0].clear_coefficients()
    assert estimate.total() == before


def test_coefficient_on_commercial_price(catalog):
    estimate = Estimate(vat_pct=0)
    work = estimate.add(Position.commercial("Работа", 10, "м2", 100))
    material = estimate.add(Position.commercial("Плитка", 10, "м2", 100, kind=MATERIAL))
    work.apply_coefficient("Ночь", "1.2")
    material.apply_coefficient("Ночь", "1.2")
    assert work.total(estimate.indices) == money(1200)
    assert material.total(estimate.indices) == money(1000)  # к материалам К по ОТ не применяется


def test_limited_costs_between_positions_and_vat(catalog):
    estimate = make(catalog)
    smr = estimate.totals().smr
    estimate.limited = [{"name": "ВЗиС", "pct": "1.8"}, {"name": "Непредвиденные", "pct": "2"}]
    totals = estimate.totals()
    assert totals.limited_total == money(smr * Decimal("0.018")) + money(smr * Decimal("0.02"))
    assert totals.vat == money(totals.subtotal * Decimal("0.2"))
    for goal in ("1000000", "999999.99", "1234567.89", "500000.01"):
        assert abs(estimate.fit_to(goal) - Decimal(goal)) <= Decimal("0.01")
    assert Estimate.from_dict(estimate.to_dict()).total() == estimate.total()


# ---------------------------------------------------------------- progress and acts


def test_progress_acts_and_payments(catalog):
    estimate = make(catalog)
    estimate.limited = [{"name": "ВЗиС", "pct": "1.8"}]
    progress.set_done(estimate, 1, "всё")
    progress.set_done(estimate, 2, "50%")
    with pytest.raises(ValueError):
        progress.set_done(estimate, 2, "1000")
    act = progress.make_act(estimate, when=datetime.date(2026, 10, 1))
    assert act["number"] == 1 and len(act["items"]) == 2
    with pytest.raises(ValueError):
        progress.set_done(estimate, 1, "10")                # меньше закрытого актом нельзя
    with pytest.raises(ValueError):
        progress.make_act(estimate)                         # всё уже закрыто
    info = progress.progress(estimate)
    assert info.acted == Decimal(act["total"])
    assert info.done == info.acted                          # одинаковое округление
    progress.add_payment(estimate, 100000)
    assert info.debt > progress.progress(estimate).debt
    progress.set_done(estimate, 2, "всё")
    second = progress.make_act(estimate, when=datetime.date(2026, 11, 1))
    assert second["from"] == "2026-10-02"
    year, since_start = progress.year_to_date(estimate, second)
    assert year == since_start == money(Decimal(act["total"]) + Decimal(second["total"]))


def test_schedule_states(catalog):
    estimate = make(catalog)
    progress.set_schedule(estimate, "30/70", start=datetime.date(2026, 1, 1))
    info = progress.progress(estimate, today=datetime.date(2026, 3, 1))
    assert [row["state"] for row in info.schedule] == ["overdue", "overdue"]
    progress.add_payment(estimate, info.schedule[0]["amount"])
    info = progress.progress(estimate, today=datetime.date(2026, 3, 1))
    assert info.schedule[0]["state"] == "paid" and info.schedule[1]["state"] == "overdue"


# ---------------------------------------------------------------- obmer and templates


def test_room_parsing_and_areas():
    room = obmer.parse_room("кухня 4,2х3,1 2,7 д1 о2")
    assert room.name == "Кухня"
    assert room.floor == Decimal("13.02") and room.perimeter == Decimal("14.60")
    assert room.walls == money(Decimal("14.6") * Decimal("2.7") - Decimal("0.9") * Decimal("2.1") - 2 * Decimal("2.25"))
    with pytest.raises(ValueError):
        obmer.parse_room("кухня большая")
    values = obmer.summary([room, obmer.parse_room("5 4 2.7 д2 о0")])
    assert values["S"] == Decimal("33.02") and values["D"] == 3 and values["R"] == 2


def test_safe_eval_refuses_code():
    assert templates.safe_eval("ceil(S * 0.05)", {"S": Decimal(56)}) == 3
    for bad in ("__import__('os')", "S.__class__", "open('x')"):
        with pytest.raises((ValueError, SyntaxError)):
            templates.safe_eval(bad, {"S": Decimal(1)})


def test_template_builds_matched_estimate(catalog, offers):
    values, notes = templates.variables(56, 2.7, rooms=2, baths=1)
    assert any("периметр" in note for note in notes)
    matches = templates.build_template("capital", values, catalog, offers)
    assert len(matches) >= 18
    assert all(match.status == "matched" for match in matches)
    exact, notes = templates.variables(56, 2.7, rooms=2, baths=1, perimeter=60, doors=3, windows=2)
    assert not notes and exact["W"] == money(60 * Decimal("2.7") - 3 * Decimal("1.89") - 2 * Decimal("2.25"))


# ---------------------------------------------------------------- materials by norms


def test_materials_by_norms(catalog, offers):
    estimate = make(catalog, "штукатурка стен 120 м2\nшпаклевка стен 120 м2\nокраска стен 120 м2\nдемонтаж штукатурки 50 м2")
    added = consumption.apply(estimate, offers)
    names = {p.name: p for p in added}
    plaster = names["Штукатурка гипсовая 30 кг"]
    assert plaster.qty == 36                                 # 120 × 8,5 × 1,05 / 30 = 35,7 → 36 мешков
    assert plaster.unit_price == 650 and plaster.unit_cost is not None
    assert "Краска интерьерная 9 л" in names
    assert all("Демонтаж" not in p.note for p in added)
    consumption.apply(estimate, offers)
    assert sum(1 for p in estimate.positions if p.source == consumption.SOURCE) == len(added)   # без дублей


# ---------------------------------------------------------------- audit


def foreign_estimate() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Смета"])
    sheet.append(["№", "Наименование работ", "Ед. изм.", "Кол-во", "Цена, руб.", "Сумма, руб."])
    sheet.append([1, "Штукатурка стен по маякам", "м2", 120, 2500, 300000])
    sheet.append([2, "Шпатлёвка стен", "м2", 120, 800, 99000])
    sheet.append([3, "Покраска стен", "м2", 120, 650, 78000])
    sheet.append([4, "Покраска стен", "м2", 120, 650, 78000])
    sheet.append([None, "Итого", None, None, None, 560000])
    sheet.append([None, "НДС 20%", None, None, None, 100000])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_audit_finds_measurable_problems(catalog, tmp_path):
    report = audit.audit(foreign_estimate(), "чужая.xlsx", catalog, catalog.index())
    issues = "\n".join(issue for row in report.rows for issue in row.issues)
    assert "арифметика" in issues
    assert "повтор строки" in issues
    assert "дороже справочника" in issues
    assert any("Итого" in item for item in report.totals_issues)
    assert any("НДС" in item for item in report.totals_issues)
    assert any("грунтование" in item for item in report.missing)
    assert report.overprice > 0
    path = audit.report_xlsx(report, tmp_path / "r.xlsx")
    assert load_workbook(path).active.max_row > 5


def test_audit_rejects_non_estimate(catalog):
    with pytest.raises(ValueError):
        audit.audit(b"a;b\n1;2\n", "x.csv", catalog, catalog.index())


# ---------------------------------------------------------------- versions and reprice


def test_diff_and_additional_agreement(catalog, tmp_path):
    old = make(catalog)
    new = old.copy()
    new.positions[0].qty = Decimal(150)
    for match in build_from_text("плинтус 40 м", catalog):
        new.add(match.position)
    new.remove(2)
    change = versions.diff(old, new)
    assert len(change.added) == 1 and len(change.removed) == 1 and len(change.changed) == 1
    assert change.delta == new.total() - old.total()
    path = docs.ds_xlsx(new, change, 1, tmp_path / "ds.xlsx")
    text = " ".join(str(c) for row in load_workbook(path).active.iter_rows(values_only=True) for c in row if c)
    assert "ДОПОЛНИТЕЛЬНОЕ СОГЛАШЕНИЕ" in text and "исключаются" in text


def test_reprice_by_index_and_offers(catalog, offers):
    estimate = make(catalog)
    estimate.positions[2].unit_price = Decimal(1000)
    result = reprice.reprice(estimate, indices=catalog.indices["ПРИМЕР Москва IV кв. 2026"])
    assert result.new_total > result.old_total and result.estimate is not estimate
    assert estimate.indices.ot == 1                                   # оригинал не тронут
    by_price = reprice.reprice(estimate, offers=offers)
    assert by_price.estimate.positions[2].unit_price == 1550
    assert "Уведомляем" in by_price.letter(estimate)


# ---------------------------------------------------------------- documents


def test_ks2_ks3_naryad_kp(catalog, tmp_path):
    estimate = make(catalog, contract="№ 1 от 01.10.2026")
    progress.set_done(estimate, 1, "всё")
    act = progress.make_act(estimate)
    ks2 = " ".join(str(c) for r in load_workbook(docs.ks2_xlsx(estimate, act, tmp_path / "a.xlsx")).active.iter_rows(values_only=True) for c in r if c)
    assert "АКТ № 1" in ks2 and "Всего по акту" in ks2
    ks3 = " ".join(str(c) for r in load_workbook(docs.ks3_xlsx(estimate, act, tmp_path / "b.xlsx")).active.iter_rows(values_only=True) for c in r if c)
    assert "СПРАВКА № 1" in ks3
    estimate.positions[0].crew_price = Decimal(500)
    rows, total, missing = docs.crew_rows(estimate)
    assert total == money(500 * 120) and missing == 1
    docs.naryad_xlsx(estimate, tmp_path / "n.xlsx")
    groups = docs.kp_groups(estimate)
    assert sum(value for _, value in groups) == estimate.total()
    pdf = docs.kp_pdf(estimate, tmp_path / "kp.pdf")
    assert pdf.read_bytes().startswith(b"%PDF")


# ---------------------------------------------------------------- funnel and billing


def test_funnel_stats():
    won, lost, draft = Estimate(), Estimate(), Estimate()
    for item in (won, lost, draft):
        item.add(Position.commercial("x", 1, "шт", 100000))
    funnel.set_status(won, "sent")
    funnel.set_status(won, "agreed")
    funnel.set_status(lost, "sent")
    funnel.set_status(lost, "lost", "дорого")
    stats = funnel.stats([won, lost, draft])
    assert stats.conversion == Decimal("50.0")
    assert stats.lost_reasons == {"дорого": 1}
    assert stats.in_work == won.total()


def test_billing_limits_promo_and_idempotent_payments(tmp_path, monkeypatch):
    monkeypatch.setenv("SMETA_FREE_ESTIMATES", "2")
    monkeypatch.setenv("SMETA_PROMO", "START:14")
    monkeypatch.delenv("SMETA_ADMIN_IDS", raising=False)
    monkeypatch.delenv("ADMIN_IDS", raising=False)
    store = Store(tmp_path / "b.db")
    assert billing.can_create(store, 5)[0]
    store.save_estimate(5, Estimate())
    store.save_estimate(5, Estimate())
    assert not billing.can_create(store, 5)[0]
    until = billing.activate(store, 5, "sub:month", 99000, "RUB", "charge-1")
    assert billing.is_active(store, 5) and billing.can_create(store, 5)[0]
    assert billing.activate(store, 5, "sub:month", 99000, "RUB", "charge-1") == until    # повтор не продлевает
    ok, text = billing.redeem(store, 5, "start")
    assert ok and "+14" in text
    assert not billing.redeem(store, 5, "START")[0]
    monkeypatch.setenv("SMETA_PAYMENT_TOKEN", "123:TEST")
    params = billing.invoice(billing.plan("year"))
    assert params["currency"] == "RUB" and params["prices"][0].amount == 790000
    monkeypatch.setenv("SMETA_PAYMENT_RECEIPT", "1")
    assert "receipt" in billing.invoice(billing.plan("year"))["provider_data"]
    monkeypatch.setenv("SMETA_ADMIN_IDS", "6")
    assert billing.is_active(store, 6)
    store.close()


def test_expiring_subscription_reminder_once(tmp_path):
    store = Store(tmp_path / "e.db")
    now = datetime.datetime.now(datetime.timezone.utc)
    store.extend_subscription(9, 2, now=now)
    assert len(store.expiring_subscriptions(now)) == 1
    store.update_settings(9, sub_reminded=store.user(9)["paid_until"])
    assert not store.expiring_subscriptions(now)
    store.close()


# ---------------------------------------------------------------- reverse tax pricing


@pytest.mark.parametrize("regime,vat", [(tax.USN6, 0), (tax.USN15, 0), (tax.OSNO_OOO, 20), (tax.OSNO_IP, 20)])
def test_price_for_hits_the_target(regime, vat):
    revenue, result, inputs = tax.price_for(250000, regime, materials=180000, payroll=120000, vat_pct=vat,
                                            ip=regime == tax.OSNO_IP)
    assert result.net(inputs) >= Decimal(250000)
    smaller = tax.price_for(250000, regime, materials=180000, payroll=120000, vat_pct=vat,
                            ip=regime == tax.OSNO_IP)[0] - Decimal("0.02")
    assert smaller < revenue
