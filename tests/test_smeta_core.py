"""Core estimate logic: parsing, matching, money math, cuts, exports, margin."""
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from smeta.builder import FROM_PRICE_LIST, MATCHED, NEEDS_PRICE, PRICED, UNIT_MISMATCH, build_from_text
from smeta.catalog import Catalog, Indices
from smeta.export_pdf import estimate_to_pdf
from smeta.export_xlsx import estimate_to_xlsx, margin_to_xlsx
from smeta.materials import OfferBook
from smeta.model import MATERIAL, Company, Estimate, Position, inn_is_valid
from smeta.money import money, qty_str, rub
from smeta.parser import parse_text
from smeta.profit import build_report
from smeta.store import Store
from smeta.units import norm_unit, parse_unit

RATES_CSV = """Шифр;Наименование;Ед.изм.;ОТ;ЭМ;ЗПМ;МАТ;ТЗ;Вид работ;Ключевые слова
Т-01;Штукатурка стен;100 м2;45000;2500;700;0;110;отделочные;оштукатурить выравнивание стен
Т-02;Демонтаж штукатурки;100 м2;25000;0;0;0;62;демонтаж;сбить снять штукатурка
Т-03;Установка унитаза;шт;4000;0;0;0;9;сантехнические;унитаз
"""
NRSP_CSV = "work_type;nr;sp\nотделочные;105;55\nдемонтаж;80;40\nсантехнические;100;60\nпрочие;95;50\n"
OFFERS_CSV = (
    "material;unit;price;discount;retail;supplier;payment\n"
    "Керамогранит 600х600;м2;1000;25;1500;ТД Плитка;безнал\n"
    "Кабель ВВГнг 3х2.5;м;100;20;140;ЭлектроОпт;нал\n"
)


@pytest.fixture
def catalog() -> Catalog:
    cat = Catalog.from_csv(RATES_CSV, source="тест")
    cat.nr_sp.update(Catalog(nr_sp={}).nr_sp)
    from smeta.catalog import parse_nr_sp_csv

    cat.nr_sp.update(parse_nr_sp_csv(NRSP_CSV))
    cat.indices["тест"] = Indices(region="тест", ot=Decimal("1"), em=Decimal("1"), mat=Decimal("1"))
    return cat


@pytest.fixture
def offers() -> OfferBook:
    return OfferBook.from_csv(OFFERS_CSV)


# ---------- helpers ----------


def test_money_formats():
    assert rub("1234567.891") == "1 234 567,89"
    assert rub(Decimal("-100")) == "-100,00"
    assert money("1 234,56") == Decimal("1234.56")
    assert qty_str(Decimal("12.500")) == "12,5"
    assert qty_str(3) == "3"


def test_units():
    assert norm_unit("кв. м") == "м2"
    assert norm_unit("пог.м") == "м"
    assert parse_unit("100 м2") == (Decimal(100), "м2")
    assert parse_unit("шт") == (Decimal(1), "шт")


def test_inn_checksum():
    assert inn_is_valid("7707083893")
    assert not inn_is_valid("1234567890")
    assert not inn_is_valid("770708389")


# ---------- parser ----------


@pytest.mark.parametrize(
    "line,name,qty,unit,price",
    [
        ("Штукатурка стен 120 м2", "Штукатурка стен", "120", "м2", None),
        ("стяжка пола 85 м2 по 700", "стяжка пола", "85", "м2", "700"),
        ("плитка 2,7х4,2 м2 х 1500", "плитка", "11.34", "м2", "1500"),
        ("вывоз мусора 3 т цена 2 500", "вывоз мусора", "3", "т", "2500"),
        ("мат: керамогранит 22 м2 1550", "керамогранит", "22", "м2", "1550"),
        ("мат: штукатурка гипсовая 30 кг 170 шт 650", "штукатурка гипсовая 30 кг", "170", "шт", "650"),
        ("мат: кабель ввг 3х2.5 180 м 135", "кабель ввг 3х2.5", "180", "м", "135"),
        ("шпаклевка 300м2 350 рублей", "шпаклевка", "300", "м2", "350"),
        ("розетки 24 точки", "розетки", "24", "точка", None),
        ("ламинат 56,5 кв м", "ламинат", "56.5", "м2", None),
    ],
)
def test_parse_line(line, name, qty, unit, price):
    draft = parse_text(line)[0]
    assert draft.name == name
    assert draft.qty == Decimal(qty)
    assert draft.unit == unit
    assert draft.price == (Decimal(price) if price else None)


def test_parse_splits_and_skips_junk():
    drafts = parse_text("штукатурка 10 м2; покраска 10 м2\n\n42\n")
    assert [d.name for d in drafts] == ["штукатурка", "покраска"]


def test_material_prefix_sets_kind():
    assert parse_text("мат: клей 10 шт")[0].kind == MATERIAL
    assert parse_text("штукатурка 10 м2")[0].kind == "work"


# ---------- catalog ----------


def test_search_does_not_confuse_demolition(catalog):
    rate, score = catalog.best("штукатурка стен")
    assert rate.code == "Т-01"
    rate, _ = catalog.best("демонтаж штукатурки")
    assert rate.code == "Т-02"


def test_search_by_code(catalog):
    assert catalog.search("Т-03")[0][0].name == "Установка унитаза"


def test_search_misses_unknown_work(catalog):
    assert catalog.best("подшив сайдингом") is None


def test_import_accepts_alternative_headers():
    cat = Catalog.from_csv("code,name,unit,ot\nX-1,Моя работа,шт,100\n")
    assert cat.rates[0].code == "X-1"
    with pytest.raises(ValueError):
        Catalog.from_csv("a;b\n1;2\n")


# ---------- model math ----------


def test_normative_position_math(catalog):
    rate = catalog.by_code("Т-01")
    position = Position.from_rate(rate, 120, catalog.norms(rate.work_type))
    costs = position.costs(catalog.index("тест"))
    # 120 m2 of a "100 m2" rate -> factor 1.2
    assert costs.ot == money("54000")
    assert costs.em == money("3000")
    assert costs.nr == money((54000 + 840) * Decimal("1.05"))
    assert costs.sp == money((54000 + 840) * Decimal("0.55"))
    assert position.total(catalog.index("тест")) == money("144744")


def test_indices_convert_base_prices(catalog):
    rate = catalog.by_code("Т-01")
    position = Position.from_rate(rate, 100, catalog.norms(rate.work_type))
    plain = position.total(Indices())
    doubled = position.total(Indices(ot=Decimal(2), em=Decimal(2), mat=Decimal(2)))
    assert doubled == money(plain * 2)


def test_vat_and_totals(catalog):
    estimate = Estimate(indices=catalog.index("тест"), vat_pct=Decimal(20))
    estimate.add(Position.commercial("Работа", 2, "шт", 1000))
    totals = estimate.totals()
    assert totals.subtotal == money(2000)
    assert totals.vat == money(400)
    assert totals.total == money(2400)
    estimate.vat_pct = Decimal(0)
    assert estimate.total() == money(2000)


def test_cut_and_scope(catalog):
    estimate = Estimate(indices=catalog.index("тест"))
    estimate.add(Position.commercial("Работа", 1, "шт", 1000))
    estimate.add(Position.commercial("Материал", 1, "шт", 1000, kind=MATERIAL))
    estimate.vat_pct = Decimal(0)
    assert estimate.adjust(-20) == money(1600)
    assert estimate.discount_pct() == Decimal("-20.0")
    estimate.reset_adjustments()
    estimate.adjust(-50, MATERIAL)
    assert estimate.total() == money(1500)
    with pytest.raises(ValueError):
        estimate.adjust(-100)


def test_fit_to_hits_the_target_exactly(catalog):
    estimate = Estimate(indices=catalog.index("тест"))
    for price in (1234, 5678, 999):
        estimate.add(Position.commercial("Работа", 3, "м2", price))
    assert estimate.fit_to(900000) == money(900000)
    assert estimate.fit_to("1000000.55") == money("1000000.55")


def test_cut_does_not_cut_the_self_cost(catalog):
    rate = catalog.by_code("Т-01")
    estimate = Estimate(indices=catalog.index("тест"))
    estimate.add(Position.from_rate(rate, 120, catalog.norms(rate.work_type)))
    floor_before = estimate.floor()
    estimate.adjust(-40)
    assert estimate.floor() == floor_before
    assert any("ниже себестоимости" in w for w in estimate.warnings())


def test_floor_without_overheads(catalog):
    rate = catalog.by_code("Т-01")
    estimate = Estimate(indices=catalog.index("тест"), floor_with_nr=False)
    estimate.add(Position.from_rate(rate, 120, catalog.norms(rate.work_type)))
    assert estimate.floor() == money("57000")  # 54000 ОТ + 3000 ЭМ


def test_warnings_check_header_and_prices(catalog):
    estimate = Estimate(indices=catalog.index("тест"), company=Company(name="ООО", inn="1234567890"))
    estimate.add(Position.commercial("Без цены", 1, "шт", 0))
    joined = " ".join(estimate.warnings())
    assert "контрольной суммы" in joined
    assert "цена не задана" in joined


def test_serialisation_roundtrip(catalog):
    estimate = Estimate(number="СМ-1", indices=catalog.index("тест"))
    rate = catalog.by_code("Т-01")
    estimate.add(Position.from_rate(rate, 120, catalog.norms(rate.work_type)))
    estimate.adjust(-15)
    clone = Estimate.from_dict(estimate.to_dict())
    assert clone.total() == estimate.total()
    assert clone.history == estimate.history


# ---------- builder ----------


def test_builder_statuses(catalog, offers):
    matches = build_from_text(
        "штукатурка стен 120 м2\n"
        "подшив сайдингом 30 м2\n"
        "подшив сайдингом 30 м2 по 900\n"
        "установка унитаза 120 м2\n"
        "мат: керамогранит 22 м2\n",
        catalog,
        offers=offers,
    )
    assert [m.status for m in matches] == [MATCHED, NEEDS_PRICE, PRICED, UNIT_MISMATCH, FROM_PRICE_LIST]
    assert matches[1].position.total(catalog.index("тест")) == 0
    assert matches[2].position.unit_price == Decimal(900)
    assert matches[4].position.unit_price == Decimal(1500)      # розница заказчику
    assert matches[4].position.unit_cost == Decimal(750)        # прайс 1000 − скидка 25 %


def test_customer_price_beats_the_catalog(catalog):
    match = build_from_text("штукатурка стен 100 м2 по 500", catalog)[0]
    assert match.status == PRICED
    assert match.position.total(catalog.index("тест")) == money(50000)
    assert match.position.code == "Т-01"


# ---------- profit ----------


def test_profit_report_splits_works_and_materials(catalog, offers):
    estimate = Estimate(indices=catalog.index("тест"), vat_pct=Decimal(0))
    rate = catalog.by_code("Т-01")
    estimate.add(Position.from_rate(rate, 100, catalog.norms(rate.work_type)))
    estimate.add(Position.commercial("Керамогранит 600х600", 10, "м2", 1500, kind=MATERIAL))
    estimate.add(Position.commercial("Кабель ВВГнг 3х2.5", 100, "м", 140, kind=MATERIAL))
    report = build_report(estimate, offers)
    assert report.materials_margin == money(10 * (1500 - 750) + 100 * (140 - 80))
    assert report.by_payment["нал"] == money(8000)
    assert report.margin_total == money(report.works_margin + report.materials_margin)
    assert "в карман" in report.as_text()


def test_profit_flags_loss_after_deep_cut(catalog, offers):
    estimate = Estimate(indices=catalog.index("тест"))
    rate = catalog.by_code("Т-01")
    estimate.add(Position.from_rate(rate, 100, catalog.norms(rate.work_type)))
    estimate.adjust(-45)
    report = build_report(estimate, offers)
    assert report.losses
    assert report.margin_total < 0


# ---------- exports ----------


def test_xlsx_has_header_and_totals(tmp_path, catalog):
    estimate = Estimate(
        number="СМ-2026-0007",
        title="Ремонт офиса",
        customer="ООО «Ромашка»",
        company=Company(name="ООО «СтройПро»", inn="7707083893", signer="Иванов И.И."),
        indices=catalog.index("тест"),
    )
    rate = catalog.by_code("Т-01")
    estimate.add(Position.from_rate(rate, 120, catalog.norms(rate.work_type)))
    path = estimate_to_xlsx(estimate, tmp_path / "s.xlsx")
    text = "\n".join(
        str(cell) for row in load_workbook(path)["Смета"].iter_rows(values_only=True) for cell in row if cell
    )
    assert "ООО «СтройПро»" in text
    assert "ИНН 7707083893" in text
    assert "ЛОКАЛЬНАЯ СМЕТА № СМ-2026-0007" in text
    assert "ВСЕГО по смете" in text
    assert "Т-01" in text
    assert str(float(estimate.total())) in text


def test_xlsx_never_leaks_purchase_prices(tmp_path, catalog, offers):
    estimate = Estimate(indices=catalog.index("тест"))
    estimate.add(Position.commercial("Керамогранит 600х600", 10, "м2", 1500, kind=MATERIAL, cost=750))
    path = estimate_to_xlsx(estimate, tmp_path / "s.xlsx")
    values = [
        cell for row in load_workbook(path)["Смета"].iter_rows(values_only=True) for cell in row if cell is not None
    ]
    assert 750 not in values and "750" not in [str(v) for v in values]
    assert not any("скидк" in str(v).lower() or "поставщ" in str(v).lower() for v in values)


def test_margin_file_is_marked_internal(tmp_path, catalog, offers):
    estimate = Estimate(indices=catalog.index("тест"))
    estimate.add(Position.commercial("Керамогранит 600х600", 10, "м2", 1500, kind=MATERIAL))
    report = build_report(estimate, offers)
    path = margin_to_xlsx(report, tmp_path / "m.xlsx", estimate)
    text = "\n".join(
        str(cell) for row in load_workbook(path)["Профит"].iter_rows(values_only=True) for cell in row if cell
    )
    assert "НЕ ОТПРАВЛЯТЬ ЗАКАЗЧИКУ" in text
    assert "ТД Плитка" in text


def test_pdf_is_written(tmp_path, catalog):
    estimate = Estimate(number="СМ-1", company=Company(name="ООО «СтройПро»", inn="7707083893"),
                        indices=catalog.index("тест"))
    estimate.add(Position.commercial("Работа", 2, "шт", 1000))
    path = estimate_to_pdf(estimate, tmp_path / "s.pdf")
    assert path.exists() and path.stat().st_size > 1000
    assert path.read_bytes().startswith(b"%PDF")


# ---------- store ----------


def test_store_roundtrip(tmp_path, catalog):
    store = Store(tmp_path / "db.sqlite")
    store.save_company(1, Company(name="ООО «СтройПро»", inn="7707083893"))
    assert store.company(1).name == "ООО «СтройПро»"
    assert store.next_number(1).endswith("-0001")
    assert store.next_number(1).endswith("-0002")

    estimate = Estimate(number="СМ-1", indices=catalog.index("тест"))
    estimate.add(Position.commercial("Работа", 2, "шт", 1000))
    estimate_id = store.save_estimate(1, estimate)
    estimate.add(Position.commercial("Ещё", 1, "шт", 500))
    store.save_estimate(1, estimate, estimate_id)
    assert len(store.load_estimate(1, estimate_id).positions) == 2
    assert store.load_estimate(2, estimate_id) is None          # чужую смету не отдаём
    assert len(store.list_estimates(1)) == 1

    store.add_import(1, "rates", "my.csv", RATES_CSV, 3)
    merged = store.catalog(1, catalog)
    assert merged.by_code("Т-01") is not None
    store.add_import(1, "offers", "p.csv", OFFERS_CSV, 2)
    assert store.offers(1).best("керамогранит").supplier == "ТД Плитка"
    store.update_settings(1, vat_pct=Decimal(0), current_id=estimate_id)
    assert store.settings(1)["vat_pct"] == 0
    assert store.settings(1)["current_id"] == estimate_id
    store.close()
