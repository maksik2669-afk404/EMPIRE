#!/usr/bin/env python3
"""Builds the WINTER ARC tracker as .xlsx (same design as Code.gs) for direct upload to Google Drive.

Google converts it to a Google Sheet. After the upload: select C9:J100 on the HIM and HER sheets ->
Insert -> Checkbox; every formula already counts TRUE values.
"""
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

YEAR, DAYS, T = 2026, 92, 0.8
FIRST, LAST, N = 9, 9 + 92 - 1, 8
C = dict(bg="0A1020", panel="111A2E", panel2="15213A", band="0D1528", text="E6F0FF", muted="7F95BA", dim="4A5B7C")
PROFILES = [
    dict(sheet="❄ HIM", title="WINTER ARC · HIM", accent="7DD3FC", strong="38BDF8", checked="0C3553", today="123E63",
         habits=["🌅 Подъём до 6:30", "🏋️ Тренировка", "🚿 Холодный душ", "📖 Чтение 20 мин", "💧 2 л воды",
                 "🥩 Белок и чистое питание", "📵 Без соцсетей до 12:00", "🎯 1 час на цель"]),
    dict(sheet="❄ HER", title="WINTER ARC · HER", accent="C7D2FE", strong="A5B4FC", checked="232A57", today="2C3470",
         habits=["🌅 Подъём до 7:00", "🧘 Тренировка / пилатес", "🧴 Уход утром и вечером", "📖 Чтение 20 мин",
                 "💧 2 л воды", "🥗 Чистое питание", "📓 3 благодарности", "🚶 10 000 шагов"]),
]


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def px(width_px):  # pixels -> Excel column width units
    return round((width_px - 5) / 7, 2)


def paint(ws, rows, cols, color):
    f = fill(color)
    for row in ws.iter_rows(min_row=1, max_row=rows, min_col=1, max_col=cols):
        for cell in row:
            cell.fill = f


def style(ws, ref, *, font=None, bg=None, align=None, fmt=None):
    rows = ws[ref] if ":" in ref else ((ws[ref],),)
    for row in rows:
        for cell in row:
            if font:
                cell.font = font
            if bg:
                cell.fill = fill(bg)
            if align:
                cell.alignment = align
            if fmt:
                cell.number_format = fmt


BOLD = Font(name="Montserrat", size=10, bold=True, color=C["text"])
MID = Alignment(horizontal="center", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
GAP = Side(style="thick", color=C["bg"])


def build_start(ws):
    ws.title = "❄ СТАРТ"
    ws.sheet_properties.tabColor = "E0F2FE"
    ws.sheet_view.showGridLines = False
    paint(ws, 26, 8, C["bg"])
    ws.column_dimensions["A"].width = px(28)
    for col in "BCDEFG":
        ws.column_dimensions[col].width = px(130)
    lines = [
        (2, f"❄  WINTER ARC {YEAR}", Font(name="Oswald", size=30, bold=True, color=C["text"]), 45),
        (3, f"{DAYS} дней, чтобы к Новому году стать версией себя, которой будешь гордиться.",
         Font(name="Montserrat", size=11, color="7DD3FC"), None),
        (5, "КАК ПОЛЬЗОВАТЬСЯ", Font(name="Montserrat", size=9, bold=True, color=C["muted"]), None),
        (6, "1.  Открой свой лист — ❄ HIM или ❄ HER (вкладки внизу).", None, None),
        (7, "2.  Впиши свои привычки в заголовки таблицы (строка 8) — всё пересчитается само.", None, None),
        (8, "3.  Каждый вечер отмечай галочки. Сегодняшний день подсвечен.", None, None),
        (9, f"4.  День засчитан, если выполнено ≥ {int(T * 100)}% привычек. Серия — дни подряд без провала.", None, None),
        (10, f"5.  Сон, вес, настроение и заметка — по желанию. Через {DAYS} дней это твой дневник перемен.", None, None),
        (12, "ПРАВИЛА АРКИ", Font(name="Montserrat", size=9, bold=True, color=C["muted"]), None),
        (13, "—  Никаких нулевых дней: в плохой день сделай хотя бы минимум.", None, None),
        (14, "—  Не обсуждай — показывай. Результат скажет за тебя.", None, None),
        (15, "—  Сорвался — не жди понедельника. Начни со следующей галочки.", None, None),
        (16, "—  Меньше экрана, больше жизни. Зима — время строить.", None, None),
        (18, "МОИ 3 ЦЕЛИ К 31.12", Font(name="Montserrat", size=9, bold=True, color=C["muted"]), None),
        (23, f"СТАРТ 01.10.{YEAR}   ·   ФИНИШ 31.12.{YEAR}", Font(name="Montserrat", size=9, color=C["dim"]), None),
    ]
    for r, text, font, height in lines:
        ws.merge_cells(f"B{r}:G{r}")
        ws[f"B{r}"] = text
        ws[f"B{r}"].font = font or Font(name="Montserrat", size=11, color=C["text"])
        ws[f"B{r}"].alignment = LEFT
        ws.row_dimensions[r].height = height or 20
    for i in range(3):
        r = 19 + i
        ws[f"B{r}"] = f"{i + 1}."
        ws[f"B{r}"].font = Font(name="Montserrat", size=11, bold=True, color="7DD3FC")
        ws[f"B{r}"].alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(f"C{r}:G{r}")
        style(ws, f"C{r}:G{r}", bg=C["panel2"])
        ws.row_dimensions[r].height = 24


def build_tracker(wb, p):
    ws = wb.create_sheet(p["sheet"])
    ws.sheet_properties.tabColor = p["strong"]
    ws.sheet_view.showGridLines = False
    paint(ws, LAST + 6, 18, C["bg"])
    widths = [64, 42] + [88] * N + [58, 108, 64, 64, 64, 240, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = px(w)
    for r, h in [(1, 56), (2, 26), (3, 22), (4, 44), (5, 34), (6, 10), (7, 26), (8, 58)]:
        ws.row_dimensions[r].height = h * 0.75
    for r in range(FIRST, LAST + 1):
        ws.row_dimensions[r].height = 21

    start = f"DATE({YEAR},10,1)"
    end = f"({start}+{DAYS - 1})"
    col = lambda L: f"{L}{FIRST}:{L}{LAST}"  # noqa: E731

    ws.merge_cells("A1:P1")
    ws["A1"] = f"  ❄ {p['title']} {YEAR}"
    ws["A1"].font = Font(name="Oswald", size=26, bold=True, color=C["text"])
    ws.merge_cells("A2:P2")
    ws["A2"] = f"   01.10 → 31.12  ·  {DAYS} дней  ·  день засчитан при ≥ {int(T * 100)}%  ·  не рви серию"
    ws["A2"].font = Font(name="Montserrat", size=10, color=p["accent"])

    kpis = [
        ("A", "C", "ОБЩИЙ ПРОГРЕСС", f'=IFERROR(AVERAGEIFS({col("K")},{col("A")},"<="&TODAY()),0)', "0%"),
        ("D", "F", "ДНЕЙ ЗАСЧИТАНО", f'=COUNTIF({col("K")},">="&{T})', "0"),
        ("G", "I", "СЕРИЯ СЕЙЧАС", f'=MAX(IFERROR(INDEX({col("Q")},MATCH(TODAY(),{col("A")},0)),0),'
                                  f'IFERROR(INDEX({col("Q")},MATCH(TODAY()-1,{col("A")},0)),0))', "0"),
        ("J", "L", "ЛУЧШАЯ СЕРИЯ", f'=MAX({col("Q")})', "0"),
        ("M", "P", "ДО ФИНИША", f'=IF(TODAY()<{start},"старт через "&({start}-TODAY())&" дн.",'
                               f'IF(TODAY()>{end},"арка пройдена ❄",({end}-TODAY())&" дн."))', "@"),
    ]
    for a, b, label, formula, fmt in kpis:
        ws.merge_cells(f"{a}3:{b}3")
        ws.merge_cells(f"{a}4:{b}4")
        ws[f"{a}3"] = label
        ws[f"{a}4"] = formula
        style(ws, f"{a}3:{b}4", bg=C["panel"], align=CENTER)
        ws[f"{a}3"].font = Font(name="Montserrat", size=8, bold=True, color=C["muted"])
        ws[f"{a}4"].font = Font(name="Oswald", size=20, bold=True, color=p["accent"])
        ws[f"{a}4"].number_format = fmt
        for row in ws[f"{a}3:{b}4"]:  # gaps between KPI cards
            for cell in row:
                cell.border = Border(left=GAP, right=GAP, top=GAP, bottom=GAP)
    ws.merge_cells("A5:P5")
    ws["A5"] = '=REPT("▰",ROUND(A4*40))&REPT("▱",40-ROUND(A4*40))&"   "&TEXT(A4,"0%")'
    ws["A5"].font = Font(name="Montserrat", size=13, color=p["accent"])
    ws["A5"].alignment = CENTER

    ws.merge_cells("A7:B7")
    ws["A7"] = "ИТОГ"
    ws["A7"].font = Font(name="Montserrat", size=8, bold=True, color=C["muted"])
    ws["A7"].alignment = CENTER
    for i in range(N):
        L = chr(67 + i)
        ws[f"{L}7"] = f'=IFERROR(COUNTIF({col(L)},TRUE)/COUNTIF({col("A")},"<="&TODAY()),0)'
    ws["K7"] = "=A4"
    style(ws, "C7:K7", font=Font(name="Montserrat", size=10, bold=True, color=p["accent"]), bg=C["panel"],
          align=CENTER, fmt="0%")

    head = ["ДАТА", "ДЕНЬ"] + p["habits"] + ["%", "ПРОГРЕСС", "СОН, ч", "ВЕС", "НАСТР. 1–5", "ЗАМЕТКА ДНЯ", ""]
    for i, h in enumerate(head):
        ws.cell(8, i + 1, h)
    style(ws, "A8:Q8", font=Font(name="Montserrat", size=9, bold=True, color=C["text"]), bg=C["panel2"], align=CENTER)
    style(ws, "C8:J8", font=Font(name="Montserrat", size=9, bold=True, color=p["accent"]))
    for row in ws["A8:P8"]:
        for cell in row:
            cell.border = Border(bottom=Side(style="medium", color=p["strong"]))

    for i in range(DAYS):
        r = FIRST + i
        ws[f"A{r}"] = f"={start}+{i}"
        ws[f"B{r}"] = f'=CHOOSE(WEEKDAY(A{r},2),"Пн","Вт","Ср","Чт","Пт","Сб","Вс")'
        ws[f"K{r}"] = f"=COUNTIF(C{r}:J{r},TRUE)/{N}"
        ws[f"L{r}"] = f'=REPT("▰",ROUND(K{r}*10))&REPT("▱",10-ROUND(K{r}*10))'
        ws[f"Q{r}"] = f"=IF(K{r}>={T},1,0)" if i == 0 else f"=IF(K{r}>={T},Q{r - 1}+1,0)"
        ws[f"A{r}"].number_format = "dd.mm"
        ws[f"A{r}"].font = BOLD
        ws[f"B{r}"].font = Font(name="Montserrat", size=9, color=C["muted"])
        ws[f"K{r}"].number_format = "0%"
        ws[f"K{r}"].font = BOLD
        ws[f"L{r}"].font = Font(name="Montserrat", size=10, color=p["accent"])
        ws[f"P{r}"].font = Font(name="Montserrat", size=10, color=C["muted"])
        for L in "ABCDEFGHIJKLMNO":
            ws[f"{L}{r}"].alignment = MID
        ws[f"M{r}"].number_format = ws[f"N{r}"].number_format = "0.0"

    for ref, lo, hi, msg in [("M", 0, 24, "Сколько часов спал(а): от 0 до 24"), ("N", 20, 300, "Вес в кг"),
                             ("O", 1, 5, "Настроение от 1 до 5")]:
        dv = DataValidation(type="decimal", operator="between", formula1=str(lo), formula2=str(hi),
                            showErrorMessage=True, error=msg, prompt=msg, showInputMessage=True)
        ws.add_data_validation(dv)
        dv.add(col(ref))

    cf = ws.conditional_formatting
    cf.add(f"A{FIRST}:P{LAST}", FormulaRule(formula=[f"ISODD(ROW()-{FIRST})"], fill=fill(C["band"])))
    cf.add(f"A{FIRST}:P{LAST}", FormulaRule(formula=[f"$A{FIRST}=TODAY()"], fill=fill(p["today"]),
                                            font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
    cf.add(f"C{FIRST}:J{LAST}", FormulaRule(formula=[f"C{FIRST}=TRUE"], fill=fill(p["checked"])))
    cf.add(f"A{FIRST}:B{LAST}", FormulaRule(formula=[f"$A{FIRST}>TODAY()"], font=Font(color=C["dim"])))
    cf.add(col("K"), ColorScaleRule(start_type="num", start_value=0, start_color="16203A",
                                    end_type="num", end_value=1, end_color=p["strong"]))

    ws.column_dimensions["Q"].hidden = True
    ws.freeze_panes = f"A{FIRST}"


def main(out):
    wb = Workbook()
    # workbook default font = light Montserrat, so unstyled cells need no per-cell font on the dark background
    from openpyxl.utils.indexed_list import IndexedList
    wb._fonts = IndexedList([Font(name="Montserrat", size=10, color=C["text"])])
    build_start(wb.active)
    for p in PROFILES:
        build_tracker(wb, p)
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "winter_arc_2026.xlsx")
